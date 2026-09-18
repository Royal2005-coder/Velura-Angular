# -*- coding: utf-8 -*-
"""Velura Master Sheet — UAT + BA/PM khảo sát cải tiến (Context: code + production)."""
from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.page import PageMargins
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.series import SeriesLabel

OUT = Path(__file__).resolve().parent / "VELURA_MasterSheet_UAT_BA_CaiTien.xlsx"
DOWNLOADS = Path(r"C:\Users\ADMIN\Downloads") / "VELURA_MasterSheet_UAT_BA_CaiTien.xlsx"

PROD_USER = "https://velura.royalai.dev"
PROD_ADMIN = "https://admin.royalai.dev"
STG_USER = "https://staging.velura.royalai.dev"
STG_ADMIN = "https://staging-admin.royalai.dev"

MEMBERS = ["Gia", "Khải", "Quỳnh", "Uyên", "Ninh", "Hân"]
SECONDARY = {
    "Gia": "Hân",
    "Khải": "Gia",
    "Quỳnh": "Khải",
    "Uyên": "Quỳnh",
    "Ninh": "Uyên",
    "Hân": "Ninh",
}

# fills
BLUE = PatternFill("solid", fgColor="1F4E79")
BLUE_CELL = PatternFill("solid", fgColor="D6EAF8")
GREEN = PatternFill("solid", fgColor="196F3D")
GREEN_CELL = PatternFill("solid", fgColor="D5F5E3")
ORANGE = PatternFill("solid", fgColor="B9770E")
ORANGE_CELL = PatternFill("solid", fgColor="FDEBD0")
YELLOW = PatternFill("solid", fgColor="F4D03F")
YELLOW_CELL = PatternFill("solid", fgColor="FCF3CF")
GREY = PatternFill("solid", fgColor="F4F6F7")
RED_CELL = PatternFill("solid", fgColor="FADBD8")
WHITE = PatternFill("solid", fgColor="FFFFFF")
NAVY_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1F4E79")
SECTION_FONT = Font(name="Calibri", bold=True, size=12, color="1F4E79")
BODY = Font(name="Calibri", size=10)
WRAP = Alignment(wrap_text=True, vertical="center")
THIN = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

# Who fills: SYS=code snapshot, HITL=tester, PM=BA/PM
UAT_COLUMNS = [
    ("Ma_TC", "SYS", "Mã test case ổn định, không đổi. Format {MaTinhNang}-TC##.", "U-02-TC01"),
    ("Nhom_TC", "SYS", "Happy / Negative / Edge / Empty / Permission / Mobile / DOCX. Để lọc đúng loại kiểm thử.", "Happy"),
    ("Tinh_nang_con", "SYS", "Hành động con trong cùng tab (ví dụ OTP, voucher, size).", "Đăng nhập SĐT"),
    ("Precondition", "SYS", "Điều kiện trước khi bấm. Tester phải lập đúng trạng thái này.", "Chưa login; giỏ trống"),
    ("Buoc_thao_tac", "HITL", "Viết từng bước như user thật: vào URL → bấm → nhập. Bắt buộc chi tiết, không viết 'test login'.", "Mở /auth/signin → chọn SĐT → nhập..."),
    ("Du_lieu_test", "HITL", "Account / SKU / OTP / file dùng khi chạy. Che mật khẩu. Ghi môi trường.", "SĐT 09xx (acc UAT Gia)"),
    ("Ket_qua_mong_doi", "SYS", "Expected từ CODE hiện tại + policy. Đối chiếu DOCX cũ ở cột riêng, không trộn.", "Vào trang chủ; badge giỏ cập nhật"),
    ("Ket_qua_thuc_te", "HITL", "Cái mắt thấy trên production/staging. Nếu khác expected thì Fail.", ""),
    ("Pass_Fail", "HITL", "Pass | Fail | Blocked | N/A. Blocked = không test được (mất data, 500, thiếu quyền).", ""),
    ("Loi_co_ban", "HITL", "Triệu chứng lỗi hiện trạng. 1 câu: làm gì → thấy gì. Không viết 'lỗi UI' chung chung.", ""),
    ("Muc_do_loi", "HITL", "Blocker (không mua/không admin được) / Major / Minor / Cosmetic.", ""),
    ("Screenshot", "HITL", "Link Drive/GitLab hoặc tên file. Bắt buộc khi Fail hoặc khi đề xuất UI.", ""),
    ("Moi_truong", "HITL", "Prod user/admin | Staging | Local. UAT chính = Production đang chạy.", "Prod user"),
    ("Thiet_bi", "HITL", "Desktop 1280+ / Mobile 390. Mỗi TC Happy phải có ít nhất 1 lần Mobile nếu tab public.", ""),
    ("Doi_chieu_DOCX", "HITL", "Khớp | Lệch | Thiếu trên UI | Thiếu trên DOCX | Không có DOCX. Bắt buộc đọc file quy trình cũ.", ""),
    ("Trich_DOCX_cu", "HITL", "Nguyên văn / mục / trang trong DOCX cũ. Nếu không có file, ghi đường dẫn đã tìm và 'Không có'.", ""),
    ("Ghi_chu_lech", "HITL", "DOCX nói A, UI đang B. Đây là input BA chuẩn hóa requirement.", ""),
    ("De_xuat_cai_tien", "HITL", "Cải tiến cụ thể cho TC này (không copy slogan). Gắn với requirement.", ""),
    ("Ly_do_de_xuat", "HITL", "Vì sao (pain user, lệch thị trường, lỗi, thiếu AI). Phải có lý do.", ""),
    ("Tac_dong_hien_tai", "HITL", "Nếu làm đề xuất: Không đổi / Thấp / Trung bình / Cao / Phá vỡ luồng hiện tại.", ""),
    ("Cai_tien_Wireframe", "HITL", "Khung bố cục mới: thêm/bớt vùng, thứ tự bước. Ghi 'giữ' nếu không đổi.", ""),
    ("Cai_tien_Mockup", "HITL", "Visual: token màu, spacing, component, trạng thái hover/empty/error.", ""),
    ("Cai_tien_Prototype", "HITL", "Click-through: nhánh, micro-interaction, thời điểm hiện modal/OTP.", ""),
    ("Cai_tien_Frontend", "HITL", "Angular: signal, empty/error, a11y, lazy. Không đề xuất HttpClient trong page.", ""),
    ("Cai_tien_BPMN", "HITL", "Cổng BPMN: thêm/bớt task, lane User vs Admin vs Hệ thống vs Payment.", ""),
    ("Cai_tien_Userflow", "HITL", "Lối vào tính năng giống production thật (entry, shortcut, deep link, resume).", ""),
    ("Tinh_nang_AI_moi", "HITL", "AI gắn đúng function này (không AI chung chung). Ghi input/output/human approve.", ""),
    ("Requirement_ID", "PM", "ID yêu cầu BA sau khi chốt. Để trống đến khi HITL đủ.", ""),
    ("Muc_uu_tien", "PM", "P0 P1 P2 P3 sau khi có đủ 6 người. Tester gợi ý, PM chốt.", ""),
    ("Effort", "PM", "S <1n / M 1–3n / L 1 sprint / XL nhiều team.", ""),
    ("Tester", "SYS", "Chủ trì tab. Vẫn phải tự chạy, không ủy quyền trắng.", ""),
    ("Reviewer_HITL", "SYS", "Thành viên thứ 2 xác nhận điểm và đề xuất (human in the loop).", ""),
    ("Ngay_test", "HITL", "YYYY-MM-DD lần chạy trên production.", ""),
    ("Trang_thai", "HITL", "Chưa test | Đang test | Đạt | Không đạt | Blocked | Cần review HITL.", ""),
    ("Jira_KAN", "PM", "Ticket sau khi chốt. Không tự invent KAN.", ""),
]

HITL_CRITERIA = [
    ("H1_DoiChieu_DOCX", "Đã mở DOCX/BPMN cũ và ghi lệch?", "Bắt buộc. Không 'nhớ' quy trình."),
    ("H2_UI_Production", "UI production có giống userflow thật (lối vào, resume, empty)?", "So với app fashion live, không so slide."),
    ("H3_Wireframe", "Cần wireframe mới? Vùng nào thừa/thiếu?", "Khung, không màu."),
    ("H4_Mockup", "Cần mockup? Token/spacing/component?", "Bám design token, không hex trong component."),
    ("H5_Prototype", "Cần prototype click-through cho nhánh này?", "OTP, payment, SSO, return."),
    ("H6_Frontend", "Gap Angular: empty/error/a11y/lazy/signal?", "KAN-8 empty UI là mẫu."),
    ("H7_BPMN", "Cổng quy trình thiếu lane/task/exception?", "User vs Admin vs Payment vs CSKH."),
    ("H8_Userflow", "Lối vào tính năng có tối?", "Header, chatbot, email deep link."),
    ("H9_AI", "AI nào gắn function này? Cần human approve?", "Gợi ý / chatbot / moderation / forecast."),
    ("H10_RuiRo", "Nếu không cải tiến, user/admin bị gì?", "Bỏ cuộc, sai đơn, mất tiền, sai quyền."),
    ("H11_DiemTong_1_5", "Chấm 1–5 mức 'đủ production' của tab này", "1=không dùng được, 5=đủ live chuyên nghiệp."),
    ("H12_ChapNhan_Chot", "Tester + reviewer đồng ý chốt requirement tab này?", "Có / Chưa / Cần workshop."),
]


def F(
    ma,
    sheet,
    app,
    ten,
    owner,
    route,
    url,
    api,
    workflow,
    known,
    ui_now,
    bpmn,
    extra,
):
    return {
        "ma": ma,
        "sheet": sheet,
        "app": app,
        "ten": ten,
        "owner": owner,
        "reviewer": SECONDARY[owner],
        "route": route,
        "url": url,
        "api": api,
        "workflow": workflow,
        "known": known,
        "ui_now": ui_now,
        "bpmn": bpmn,
        "extra": extra,
    }


FEATURES = [
    F("U-01", "U_Home", "User", "Trang chủ", "Khải", "/", f"{PROD_USER}/",
      "GET /api/user/products, GET /api/user/categories, GET /api/content/pages",
      "Guest/member mở / → banner hot → sản phẩm nổi bật → CTA BST/AI/chat. Modal quiz + widget chat trên shell.",
      "Search header chỉ chạy khi Enter. Icon chuông có trong SVG nhưng header không gắn notification.",
      "Logo, hotline 1900 1212, nav, search, wishlist/cart badge, account dropdown, footer.",
      "Visit → Browse → optionally Quiz/Chat → PDP",
      [("Happy", "Hero + featured", "Mở trang chủ production", "Banner và list sản phẩm load từ API, không trắng."),
       ("Empty", "API rỗng", "Giả lập/ghi nhận khi catalog lỗi", "Phải có empty/error, không vỡ layout (KAN-8).")]),
    F("U-02", "U_Auth_DangNhap", "User", "Đăng nhập", "Gia", "/auth/signin", f"{PROD_USER}/auth/signin",
      "POST /api/user/auth/signin; OTP /otp-verify; social-login Google/Facebook",
      "Chọn SĐT hoặc email + mật khẩu ≥8. Nếu otp_required thì modal OTP. Thành công lưu session, sync wishlist, về home.",
      "Tài khoản chưa verify bắt OTP. Social cần email từ provider.",
      "2 panel SĐT/email, hiện/ẩn mật khẩu, link quên MK / đăng ký, OTP modal.",
      "Guest → Sign-in → (OTP) → Authenticated home",
      [("Happy", "SĐT + password", "Acc UAT đã verify", "Vào home, dropdown hiện Tài khoản."),
       ("Negative", "Sai mật khẩu", "Password sai", "Báo lỗi, không vào session."),
       ("Permission", "OTP chưa verify", "Acc mới", "Mở OTP, không skip."),
       ("Happy", "Social Google", "Tài khoản Google có email", "Tạo/liên kết user Velura.")]),
    F("U-03", "U_Auth_DangKy", "User", "Đăng ký", "Gia", "/auth/signup", f"{PROD_USER}/auth/signup",
      "POST /api/user/auth/signup; GET check-exists; OTP verify",
      "Họ tên + SĐT bắt buộc, email optional, password ≥8 + confirm → OTP xác minh.",
      "Trùng SĐT/email phải chặn qua check-exists.",
      "Form đăng ký + OTP modal.",
      "Guest → Sign-up → OTP → Member",
      [("Happy", "Đăng ký mới", "SĐT chưa tồn tại", "Nhận OTP, verify xong login."),
       ("Negative", "Trùng SĐT", "SĐT đã có", "Không tạo trùng.")]),
    F("U-04", "U_Auth_QuenMK", "User", "Quên mật khẩu", "Gia", "/auth/forgot-password", f"{PROD_USER}/auth/forgot-password",
      "POST /api/user/auth/otp-send",
      "Nhập SĐT/email đã đăng ký → gửi OTP reset.",
      "Phân biệt user không tồn tại vs gửi thành công (không leak nếu policy cấm).",
      "Form quên mật khẩu.",
      "Forgot → OTP → Reset page",
      [("Happy", "Gửi OTP", "Acc UAT", "Có thông báo đã gửi."),
       ("Negative", "Identity lạ", "SĐT chưa đăng ký", "Hành vi đúng policy (không enumerate nếu yêu cầu).")]),
    F("U-05", "U_Auth_ResetMK", "User", "Đặt lại mật khẩu", "Gia", "/auth/reset-password", f"{PROD_USER}/auth/reset-password",
      "POST /api/user/auth/reset-password",
      "OTP + mật khẩu mới → đăng nhập lại được.",
      "Hết hạn OTP / reuse OTP.",
      "Form reset.",
      "Reset → Sign-in bằng MK mới",
      [("Happy", "Reset thành công", "OTP còn hạn", "Login được bằng MK mới."),
       ("Negative", "OTP sai", "OTP 000000", "Từ chối, không đổi MK.")]),
    F("U-06", "U_SP_DanhSach", "User", "Danh sách sản phẩm", "Khải", "/products", f"{PROD_USER}/products",
      "GET /api/user/products; GET /api/user/categories; query ?q=",
      "Header search Enter → /products?q=. Lọc/sort/category tùy UI hiện tại.",
      "Search không search-as-you-type.",
      "Grid sản phẩm, filter, ProductCard.",
      "Home/Search → PLP → PDP",
      [("Happy", "Mở all products", "Catalog có hàng", "Có list, giá VND."),
       ("Happy", "Search q", "Gõ từ khóa Enter", "URL có q, kết quả liên quan."),
       ("Empty", "q không ra", "Từ khóa rác", "Empty state, không spinner vô hạn.")]),
    F("U-07", "U_SP_ChiTiet", "User", "Chi tiết sản phẩm", "Khải", "/products/:id", f"{PROD_USER}/products",
      "GET /api/user/products/:id; POST cart; POST wishlist",
      "Chọn màu/size/SL → thêm giỏ / yêu thích. Ảnh, giá, mô tả, tồn.",
      "Hết hàng / combo / variant thiếu.",
      "Gallery, variant, CTA giỏ/wishlist.",
      "PLP → PDP → Cart/Wishlist",
      [("Happy", "Add to cart", "Variant còn hàng", "Badge giỏ +1, vào /cart thấy dòng."),
       ("Negative", "Thiếu size", "Không chọn variant", "Không add, có báo."),
       ("Happy", "Wishlist", "Member hoặc guest store", "Badge tim cập nhật.")]),
    F("U-08", "U_BoSuuTap", "User", "Bộ sưu tập", "Khải", "/collections", f"{PROD_USER}/collections",
      "GET products/categories hoặc content collections",
      "Landing BST → click vào nhóm → PLP/PDP.",
      "BST tĩnh vs data live.",
      "Cards bộ sưu tập.",
      "Nav BST → collection → PDP",
      [("Happy", "Mở BST", "Production", "Có nội dung, click được."),
       ("DOCX", "Tên BST", "So DOCX cũ", "Khớp tên/ý tưởng marketing.")]),
    F("U-09", "U_AI_GoiY", "User", "Gợi ý AI / outfit", "Ninh", "/ai/suggestions", f"{PROD_USER}/ai/suggestions",
      "GET /api/user/recommendations/style-profile; style-quiz",
      "Nếu có quiz: combo + category recommendations. Add cart từ gợi ý.",
      "Guest chưa quiz: phải hướng đi Style Quiz, không im lặng.",
      "Combo cards, lý do gợi ý, giá.",
      "Quiz → Suggestions → PDP/Cart",
      [("Happy", "Có quiz", "Member đã làm quiz", "Hiện combo/category."),
       ("Empty", "Chưa quiz", "Guest/member mới", "CTA sang /ai/style-quiz."),
       ("AI", "Lý do gợi ý", "Đọc reason", "Reason khớp body_shape/style_tags.")]),
    F("U-10", "U_AI_StyleQuiz", "User", "Style Quiz", "Ninh", "/ai/style-quiz", f"{PROD_USER}/ai/style-quiz",
      "GET/POST /api/user/style-quiz; POST migrate",
      "Shell có modal mời quiz. Làm quiz lưu profile (dáng, tag, ngân sách) → gợi ý.",
      "Guest quiz migrate khi login.",
      "Multi-step quiz + modal invitation.",
      "Invite → Quiz → Save → Suggestions",
      [("Happy", "Hoàn thành quiz", "Trả lời đủ bước", "Lưu được, sang gợi ý."),
       ("Edge", "Bỏ giữa chừng", "Đóng modal", "Không mất site; resume được hoặc rõ phải làm lại.")]),
    F("U-11", "U_AI_Chatbot", "User", "AI Stylist Chatbot", "Ninh", "/chatbot", f"{PROD_USER}/chatbot",
      "POST /api/v1/chat/sessions; /messages; GET messages; DELETE session; favorites; rate-limit",
      "Widget shell + trang /chatbot. Tool: sản phẩm, đơn, chính sách. Rate limit. Không lộ lỗi kỹ thuật thô.",
      "Nếu tool policy fail, bot không bịa policy (fallback hướng CSKH 1900 1212).",
      "Chat UI, session, favorite.",
      "Ask → tool → answer → optional PDP",
      [("Happy", "Hỏi outfit", "Câu 'đi làm văn phòng'", "Trả lời + gợi ý SP, không stack trace."),
       ("Happy", "Hỏi đổi trả", "Câu chính sách", "Lấy từ DB/tool, không bịa."),
       ("Happy", "Tra cứu đơn", "Có mã đơn/SĐT", "Gọi get_order_status hoặc hỏi thiếu data."),
       ("Negative", "Spam", "Gửi liên tục", "Rate limit, UX thân thiện.")]),
    F("U-12", "U_Blog", "User", "Tạp chí / Blog", "Ninh", "/blog", f"{PROD_USER}/blog",
      "GET /api/content/blogs",
      "List bài → click slug.",
      "Rỗng khi CMS chưa có bài.",
      "Journal list.",
      "Nav Blog → list → detail",
      [("Happy", "List blog", "Có bài", "Card + slug."),
       ("Empty", "Không bài", "Ghi nhận", "Empty state.")]),
    F("U-13", "U_Blog_ChiTiet", "User", "Chi tiết bài viết", "Ninh", "/blog/:slug", f"{PROD_USER}/blog",
      "GET /api/content/blogs/:slug",
      "Đọc bài, CTA sản phẩm nếu có.",
      "Slug sai → không 500 trần.",
      "Article layout.",
      "List → detail",
      [("Happy", "Mở 1 slug", "Bài tồn tại", "Title/body."),
       ("Negative", "Slug sai", "abc-xyz", "404/empty có hướng về list.")]),
    F("U-14", "U_VeChungToi", "User", "Về chúng tôi", "Hân", "/about", f"{PROD_USER}/about",
      "GET /api/content/pages/about (nếu dùng)",
      "Story thương hiệu, có thể tĩnh.",
      "Lệch copy vs DOCX.",
      "About page.",
      "Footer/nav → About",
      [("Happy", "Đọc about", "Production", "Nội dung thương hiệu."),
       ("DOCX", "Copy", "So DOCX", "Khớp sứ mệnh/địa chỉ.")]),
    F("U-15", "U_LienHe", "User", "Liên hệ + cửa hàng", "Hân", "/contact", f"{PROD_USER}/contact",
      "Form liên hệ / map fragment store-map. Header link HỆ THỐNG CỬA HÀNG.",
      "Hotline 1900 1212. Map cửa hàng. Form gửi (nếu có API) hoặc mailto.",
      "Xác nhận form có backend hay chỉ UI.",
      "Contact + map.",
      "Header shops → contact#store-map",
      [("Happy", "Mở map", "Click hệ thống cửa hàng", "Tới fragment map."),
       ("Happy", "Gửi form", "Điền đủ", "Success hoặc nêu rõ chưa có API.")]),
    F("U-16", "U_UuDai", "User", "Ưu đãi tháng", "Quỳnh", "/offers", f"{PROD_USER}/offers",
      "GET /api/user/offers; vouchers list",
      "Hiển thị KM/voucher đang active → copy mã → checkout apply.",
      "Voucher hết hạn / min order.",
      "Offers landing.",
      "Offers → copy code → checkout",
      [("Happy", "Xem ưu đãi", "Có KM active", "List mã/điều kiện."),
       ("Happy", "Copy mã", "Mã còn hạn", "Dùng được ở checkout.")]),
    F("U-17", "U_ChinhSach", "User", "Chính sách & điều khoản", "Hân", "/policies", f"{PROD_USER}/policies",
      "GET /api/content/policies; GET /api/content/policies/:id",
      "Nguồn đúng cho chatbot. Đổi trả 48h sau 'Đã giao'; hàng nguyên tem; hoàn 4–5 ngày làm việc (policy bot).",
      "Copy tĩnh lệch DB.",
      "Policy list/detail.",
      "Footer → policies → đọc",
      [("Happy", "Đọc đổi trả", "Mở policy return", "Nêu 48h / tem mác."),
       ("DOCX", "So DOCX", "In nguyên văn DOCX", "Khớp hay lệch DB.")]),
    F("U-18", "U_Wishlist", "User", "Sản phẩm yêu thích", "Khải", "/wishlist", f"{PROD_USER}/wishlist",
      "GET/POST /api/user/wishlist (member); store local guest",
      "Tim trên header. Add từ PDP. Login migrate.",
      "KAN-8 empty/error UI. Guest vs member sync.",
      "Grid wishlist, badge.",
      "PDP → wishlist → cart",
      [("Happy", "Xem list", "Đã tim ≥1 SP", "Đúng sản phẩm."),
       ("Empty", "Chưa tim", "Wishlist trống", "Empty state (KAN-8)."),
       ("Happy", "Sang giỏ", "Còn hàng", "Add cart từ wishlist.")]),
    F("U-19", "U_GioHang", "User", "Giỏ hàng", "Quỳnh", "/cart", f"{PROD_USER}/cart",
      "GET/POST /api/user/cart",
      "Sửa SL, xóa dòng, CTA checkout. Badge header.",
      "KAN-8 empty/error. Hết hàng khi tăng SL.",
      "Cart lines, totals.",
      "PDP → cart → checkout/shipping",
      [("Happy", "Sửa SL", "Giỏ có hàng", "Thành tiền đổi."),
       ("Empty", "Giỏ trống", "Xóa hết", "Empty + CTA mua (KAN-8)."),
       ("Negative", "Overstock", "SL > tồn", "Chặn, thông báo.")]),
    F("U-20", "U_TK_HoSo", "User", "Tài khoản / hồ sơ", "Uyên", "/account/profile", f"{PROD_USER}/account/profile",
      "GET/PATCH /api/user/profile; PATCH /api/user/addresses",
      "Sửa tên, SĐT, địa chỉ mặc định. Dropdown 'Tài khoản' chỉ khi logged in.",
      "Guard: một số account route cần auth; profile hiện không gắn authGuard trong routes — xác minh thực tế.",
      "Profile form + addresses.",
      "Login → profile → save address",
      [("Happy", "Sửa địa chỉ", "Member", "PATCH thành công, reload đúng."),
       ("Permission", "Chưa login", "Guest mở URL", "Phải đẩy signin hoặc read-only rõ.")]),
    F("U-21", "U_TK_TheoDoiDH", "User", "Theo dõi đơn hàng", "Quỳnh", "/account/track", f"{PROD_USER}/account/track",
      "GET order by tracking/phone (user orders)",
      "Guest/member nhập mã + SĐT → timeline trạng thái.",
      "Mã sai / đơn người khác.",
      "Track form.",
      "Email/SMS mã → track → status",
      [("Happy", "Mã đúng", "Đơn UAT", "Hiện status."),
       ("Negative", "Mã sai", "TRACK-XXX", "Không lộ đơn người khác.")]),
    F("U-22", "U_TK_DonHang", "User", "Đơn hàng của tôi", "Uyên", "/account/orders", f"{PROD_USER}/account/orders",
      "GET /api/user/orders — authGuard",
      "Member xem list đơn → vào chi tiết.",
      "Guest bị chặn guard.",
      "Order list.",
      "Login → my orders → detail",
      [("Happy", "List đơn", "Member có đơn", "Mã, tiền, status."),
       ("Permission", "Guest", "Chưa login", "Redirect signin."),
       ("Empty", "Chưa mua", "Acc mới", "Empty state.")]),
    F("U-23", "U_TK_CTDonHang", "User", "Chi tiết đơn hàng", "Uyên", "/account/orders/:id", f"{PROD_USER}/account/orders",
      "GET /api/user/orders/:id; change-payment-method; payment-callback",
      "Items, địa chỉ, thanh toán, đổi PTTT nếu policy cho. CTA đổi trả/review khi đủ điều kiện.",
      "ID đơn người khác.",
      "Order detail.",
      "List → detail → return/review/pay",
      [("Happy", "Mở đơn mình", "id hợp lệ", "Đúng dòng hàng."),
       ("Permission", "id người khác", "UUID lạ", "403/404, không leak.")]),
    F("U-24", "U_TK_DoiTra", "User", "Yêu cầu đổi trả", "Uyên", "/account/returns", f"{PROD_USER}/account/returns",
      "GET/POST /api/user/returns; POST cancel; POST /api/user/upload/evidence — authGuard",
      "Tạo yêu cầu trong cửa sổ 48h sau delivered; upload ảnh; hủy khi Chờ xác nhận hoặc Đã duyệt hồ sơ chưa gửi hàng.",
      "Ngoài 48h / sai trạng thái.",
      "Return form + list.",
      "Delivered → return request → CSKH admin",
      [("Happy", "Tạo đổi trả", "Đơn delivered trong 48h", "Tạo được, hiện list."),
       ("Negative", "Quá hạn", "Đơn >48h", "Từ chối đúng policy."),
       ("Happy", "Hủy yêu cầu", "Status cho phép", "Hủy được."),
       ("Happy", "Upload evidence", "Ảnh", "Upload /api/user/upload/evidence.")]),
    F("U-25", "U_TK_DanhGia", "User", "Đánh giá sản phẩm", "Hân", "/account/reviews", f"{PROD_USER}/account/reviews",
      "GET/POST /api/user/reviews; POST :id/reply — authGuard",
      "Member đánh giá đơn đã giao. Admin duyệt mới hiện public. Reply thread.",
      "Review chưa duyệt không lên PDP.",
      "Review form/list.",
      "Delivered → review → admin approve → PDP",
      [("Happy", "Gửi review", "Đơn completed", "Gửi được, chờ duyệt."),
       ("Permission", "Chưa giao", "Đơn pending", "Không cho review.")]),
    F("U-26", "U_CK_VanChuyen", "User", "Checkout vận chuyển & PTTT", "Quỳnh", "/checkout/shipping", f"{PROD_USER}/checkout/shipping",
      "POST /api/user/orders (member); otp-send (guest); GET/POST vouchers/apply",
      "Địa chỉ + note. Ship standard 30k / express 50k. Freeship ≥500k. PTTT COD | MoMo | VNPay. Member: xác nhận đặt. Guest: nhận OTP.",
      "KAN-8 empty checkout. Voucher min_order / chưa tới hạn.",
      "Form địa chỉ, radio ship/pay, totals.",
      "Cart → shipping → OTP/pay → confirm",
      [("Happy", "Member COD", "Login, giỏ ≥1, <500k", "Phí 30k, đặt được."),
       ("Happy", "Freeship", "Subtotal ≥500k", "Phí 0."),
       ("Happy", "Voucher", "Mã đủ điều kiện", "Giảm đúng."),
       ("Negative", "Voucher sớm", "Mã chưa hiệu lực", "INVALID_VOUCHER."),
       ("Empty", "Giỏ trống", "Vào thẳng URL", "Không đặt được, empty (KAN-8).")]),
    F("U-27", "U_CK_OTP", "User", "Checkout OTP / thanh toán", "Quỳnh", "/checkout/otp|/payment", f"{PROD_USER}/checkout/otp",
      "POST /api/user/orders/otp-verify; payment-callback; MoMo/VNPay redirect",
      "Guest nhập OTP để đặt. Cổng MoMo/VNPay callback. Member có thể skip OTP.",
      "OTP hết hạn; thanh toán pending; đổi PTTT.",
      "OTP page (dùng chung payment).",
      "Shipping → OTP/gateway → confirm",
      [("Happy", "Guest OTP", "OTP đúng", "Tạo đơn, sang confirm."),
       ("Negative", "OTP sai", "000000", "Không tạo đơn."),
       ("Happy", "MoMo/VNPay", "Sandbox/prod", "Callback cập nhật payment.")]),
    F("U-28", "U_CK_ThanhCong", "User", "Đặt hàng thành công", "Quỳnh", "/checkout/confirm", f"{PROD_USER}/checkout/confirm",
      "Order payload trên store sau place order (tracking_code)",
      "Hiện mã đơn/tracking, PTTT, địa chỉ. CTA theo dõi / tiếp tục mua.",
      "Refresh mất state nếu chỉ giữ memory.",
      "Success page.",
      "Place order → confirm → track",
      [("Happy", "Sau đặt COD", "Vừa place order", "Có tracking_code."),
       ("Edge", "F5 trang", "Refresh", "Không mất mã hoặc hướng về my orders.")]),
    F("U-29", "U_ThongBao", "User", "Thông báo (API có, UI yếu)", "Ninh", "(chưa có route page)", f"{PROD_USER}/",
      "GET /api/user/notifications; POST read / read-all",
      "API member notifications tồn tại. Header có icon-bell trong sprite nhưng không gắn menu thông báo.",
      "Gap UX: user không thấy thông báo đơn/KM trên production header.",
      "Không có page; badge chuông thiếu.",
      "Event đơn → notify → user đọc",
      [("DOCX", "So DOCX", "DOCX có chuông?", "Nếu có: Thiếu trên UI."),
       ("Happy", "API", "Member + token, gọi GET notifications (DevTools)", "200 + list. Ghi nhận UI thiếu.")]),
    F("A-01", "A_Login", "Admin", "Đăng nhập quản trị", "Gia", "/login", f"{PROD_ADMIN}/login",
      "Admin session + Google SSO entry",
      "adminGuestGuard. Email/password hoặc Google. Không phải storefront auth.",
      "Sai origin callback Google (phải {origin}/auth/callback).",
      "Auth card login.",
      "Guest admin → login → (welcome|dashboard|change-password)",
      [("Happy", "Login admin", "Acc admin UAT", "Vào dashboard nếu isAdmin."),
       ("Permission", "Member thường", "Acc không admin", "Welcome chờ quyền, không vào shell.")]),
    F("A-02", "A_GoogleSSO", "Admin", "Google SSO callback", "Gia", "/auth/callback", f"{PROD_ADMIN}/auth/callback",
      "Supabase/Google OAuth callback",
      "Google → /auth/callback trên đúng origin admin.royalai.dev. Local phải localhost.",
      "Sai redirect URI.",
      "Trang 'Đang xác thực'.",
      "Google → callback → session",
      [("Happy", "SSO prod", "Google admin", "Vào đúng post-login."),
       ("Negative", "Sai origin", "Callback user storefront", "Fail, ghi cấu hình.")]),
    F("A-03", "A_DangKy", "Admin", "Đăng ký quản trị", "Uyên", "/register", f"{PROD_ADMIN}/register",
      "Tạo tài khoản chờ cấp quyền",
      "Register xong không vào dashboard. Chờ Super Admin duyệt role.",
      "Tự cấp admin.",
      "Register form.",
      "Register → welcome → role request",
      [("Happy", "Đăng ký mới", "Email mới", "Tới welcome, không thấy menu admin."),
       ("Permission", "Tự vào /dashboard", "Gõ URL", "Guard chặn.")]),
    F("A-04", "A_ChoQuyen", "Admin", "Chờ cấp quyền (welcome)", "Uyên", "/welcome", f"{PROD_ADMIN}/welcome",
      "adminWelcomeGuard; role requests GET/POST",
      "Thành viên chưa isAdmin. Super Admin duyệt trên A_TaiKhoan.",
      "Welcome leak menu.",
      "Màn chờ.",
      "Welcome → (duyệt) → dashboard",
      [("Happy", "Xem chờ", "Acc chưa admin", "Copy rõ đang chờ."),
       ("Happy", "Sau khi được duyệt", "Lead duyệt role", "Vào được shell.")]),
    F("A-05", "A_DoiMatKhau", "Admin", "Đổi mật khẩu admin", "Gia", "/change-password", f"{PROD_ADMIN}/change-password",
      "KAN-9: UI stub, submit KHÔNG gọi API; báo 'xử lý qua email đặt lại'",
      "adminSessionGuard. Form current/new/confirm ≥12 ký tự. Local validation only.",
      "KAN-9 bắt buộc ghi Fail so với requirement 'đổi MK qua API'.",
      "Form 3 field.",
      "OTP/login force change → API (chưa có)",
      [("Negative", "Submit form", "MK mới khớp ≥12", "Hiện thông báo email — KHÔNG đổi MK server. Ghi KAN-9."),
       ("Negative", "MK lệch", "confirm khác", "Lỗi client.")]),
    F("A-06", "A_Dashboard", "Admin", "Dashboard vận hành + kinh doanh", "Ninh", "/dashboard", f"{PROD_ADMIN}/dashboard",
      "GET dashboard summary: pendingOrders, paymentErrors, returns, tickets, lowStock, reviews; revenue, AOV, promo share; recentLogs; range day/week/month/custom",
      "Tab operations / business. RBAC canSee('dashboard').",
      "Số liệu lệch timezone.",
      "KPI cards + logs.",
      "Login → dashboard → drill-down orders/returns",
      [("Happy", "Operations", "Admin", "Số pending/low stock/returns."),
       ("Happy", "Business", "Đổi tab + range", "Revenue/AOV."),
       ("Happy", "Drill-down", "Click card", "Sang đúng module.")]),
    F("A-07", "A_TaiKhoan", "Admin", "Quản lý tài khoản + role", "Hân", "/accounts", f"{PROD_ADMIN}/accounts",
      "GET accounts; roles; lock; unlock; role; role-requests approve; account-audit-logs",
      "List user, khóa/mở, đổi role (có thể 202 approval). Duyệt request từ welcome.",
      "RBAC: không phải mọi admin được đổi Super Admin.",
      "Table accounts + requests.",
      "Request → review → lock/role → audit",
      [("Happy", "List", "Admin accounts", "Có user storefront/admin."),
       ("Happy", "Lock/unlock", "User test", "Khóa xong storefront không login."),
       ("Happy", "Duyệt role", "Request pending", "Welcome user vào được."),
       ("Permission", "Role thấp", "Acc limited", "Không thấy menu nếu !canSee.")]),
    F("A-08", "A_SanPham", "Admin", "Quản lý sản phẩm", "Khải", "/products", f"{PROD_ADMIN}/products",
      "CRUD products; categories; low-stock; import-csv preview/commit; variants; combo-items; change-status; update-stock; bulk-stock; audit-logs",
      "Tạo/sửa SP, variant màu-size, combo, CSV, đổi status, tồn kho, low-stock.",
      "CSV sai schema; combo circular.",
      "Catalog admin table/forms.",
      "Create → variant → stock → publish → user PLP",
      [("Happy", "List + low stock", "Admin", "Low-stock khớp dashboard."),
       ("Happy", "Tạo SP + variant", "SKU mới", "Hiện trên storefront."),
       ("Happy", "CSV preview", "File mẫu", "Preview rồi commit."),
       ("Happy", "Combo", "Add combo item", "PDP combo đúng.")]),
    F("A-09", "A_DonHang", "Admin", "Quản lý đơn hàng", "Uyên", "/orders", f"{PROD_ADMIN}/orders",
      "GET list/get; change-status; cancel; payments/:id/resolve; audit-logs",
      "Đổi trạng thái theo máy trạng thái. Hủy. Resolve payment lỗi. Audit.",
      "Nhảy status trái BPMN (vd completed ← pending).",
      "Order table + drawer.",
      "User place → admin status → user track",
      [("Happy", "Đổi status", "Đơn mới", "User track đổi theo."),
       ("Happy", "Cancel", "Đơn cho phép hủy", "Tồn kho/hoàn tiền đúng policy."),
       ("Happy", "Resolve payment", "Payment error", "Dashboard paymentErrors giảm."),
       ("Negative", "Nhảy status cấm", "Cố completed sớm", "API từ chối.")]),
    F("A-10", "A_DanhGia", "Admin", "Quản lý đánh giá", "Hân", "/reviews", f"{PROD_ADMIN}/reviews",
      "list/get; approve; hide; reply; escalate; audit-logs",
      "Duyệt mới ra PDP. Ẩn. Trả lời. Escalate CSKH.",
      "Approve nhầm review độc hại.",
      "Review queue.",
      "User review → admin approve → PDP",
      [("Happy", "Approve", "Review pending", "Hiện trên PDP."),
       ("Happy", "Hide", "Review đã hiện", "Ẩn khỏi PDP."),
       ("Happy", "Reply/escalate", "1 review", "User thấy reply / ticket.")]),
    F("A-11", "A_DoiTra_CSKH", "Admin", "Đổi trả & CSKH (chat, ticket)", "Uyên", "/returns", f"{PROD_ADMIN}/returns",
      "returns list/get; approve-refund; approve-exchange; reject; update-status; support-tickets; chat-sessions; agent reply; service-audit-logs",
      "4 zone: chat | returns | support | logs. Duyệt hoàn/đổi, reject, ticket, trả lời chat (human).",
      "Hoàn tiền nhầm cổng; chat AI vs agent.",
      "Service desk 4 tab.",
      "User return → admin decide → refund/exchange → user",
      [("Happy", "Approve refund", "Return pending", "Đúng cổng/số tiền."),
       ("Happy", "Approve exchange", "Return pending", "Tạo hướng đổi."),
       ("Happy", "Reject", "Hồ sơ sai", "User thấy rejected."),
       ("Happy", "Agent reply chat", "Session user chatbot", "User nhận tin human."),
       ("Happy", "Ticket CSKH", "List tickets", "Đổi status.")]),
    F("A-12", "A_Gia", "Admin", "Quản lý giá", "Khải", "/pricing", f"{PROD_ADMIN}/pricing",
      "GET pricing/history; statistics; product price updates qua products API",
      "Lịch sử giá, thống kê. Sửa giá bán/sale → user PDP.",
      "Giá âm / sale > base.",
      "Price history table.",
      "Edit price → history → storefront",
      [("Happy", "Đổi giá", "1 SKU", "PDP đổi, history có dòng."),
       ("Happy", "Statistics", "Admin", "Số liệu không NaN.")]),
    F("A-13", "A_KhuyenMai", "Admin", "Khuyến mãi & voucher", "Quỳnh", "/promotions", f"{PROD_ADMIN}/promotions",
      "promotions CRUD; activate; pause; vouchers CRUD",
      "Tạo KM/voucher, bật/tắt. User /offers + checkout apply. Dashboard promotionRevenue.",
      "Overlap KM; voucher min_order.",
      "Promo + voucher tables.",
      "Create promo → activate → user apply → revenue share",
      [("Happy", "Tạo + activate", "Mã UAT", "User apply được."),
       ("Happy", "Pause", "Mã đang chạy", "User không apply."),
       ("Happy", "Voucher min order", "Đơn nhỏ", "User bị INVALID_VOUCHER.")]),
    F("A-14", "A_NhatKy", "Admin", "Nhật ký hệ thống", "Hân", "/logs", f"{PROD_ADMIN}/logs",
      "GET /api/v1/admin/audit-logs",
      "Audit ai làm gì (lock user, đổi giá, status đơn...).",
      "Thiếu actor/ip.",
      "Log table filter.",
      "Admin action → audit → this page",
      [("Happy", "Thấy log sau lock user", "Vừa lock ở A-07", "Có dòng actor+target."),
       ("Permission", "Role thấp", "Acc !canSee logs", "Không vào menu.")]),
]


def default_cases(feat):
    rows = []
    rows.append(("Happy", "Luồng chính", feat["workflow"], "Đúng UI + API mô tả cột hiện trạng."))
    rows.append(("Empty", "Empty/error", "Cắt mạng hoặc data rỗng", "Có empty/error, không trắng/spinner vô hạn (KAN-8)."))
    rows.append(("Mobile", "Viewport 390", "Chrome device toolbar", "Nav, CTA, form dùng được, không che nút."))
    rows.append(("DOCX", "Đối chiếu DOCX cũ", "Mở file quy trình cũ của tab này", "Ghi Khớp/Lệch + trích dẫn."))
    rows.extend(feat["extra"])
    # de-dupe by (nhom, con)
    seen = set()
    out = []
    for n, c, s, e in rows:
        k = (n, c)
        if k in seen:
            continue
        seen.add(k)
        out.append((n, c, s, e))
    return out


def style_header(ws, row, cols, fill, font=NAVY_FONT):
    for i, _ in enumerate(cols, 1):
        cell = ws.cell(row, i)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = THIN
    ws.row_dimensions[row].height = 36


def write_row(ws, r, values, fill=None):
    for i, v in enumerate(values, 1):
        cell = ws.cell(r, i, v)
        cell.font = BODY
        cell.alignment = WRAP
        cell.border = THIN
        if fill:
            cell.fill = fill
    ws.row_dimensions[r].height = 48


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def add_dv(ws, formula, cells):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(cells)


def banner(ws, title, subtitle):
    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:H2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Calibri", size=11, italic=True, color="5D6D7E")
    ws["A2"].alignment = WRAP
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 40
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(0.4, 0.4, 0.6, 0.6)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:4"


def build():
    wb = Workbook()

    # ----- 00 hướng dẫn -----
    ws = wb.active
    ws.title = "00_DocTruocKhiDien"
    banner(
        ws,
        "VELURA — Master Sheet UAT + khảo sát cải tiến (BA / PM / HITL)",
        "Nguồn hiện trạng = CODE đang chạy + production. Why lịch sử = git log / MR, không phải file này. File này = yêu cầu cải tiến & UAT.",
    )
    ws["A3"] = "Màu cột"
    ws["A3"].font = SECTION_FONT
    legend = [
        ("Xanh đậm / xanh nhạt", "SYS — snapshot từ code/API/route. Sửa khi sai sự thật kỹ thuật, không xóa mã."),
        ("Xanh lá", "HITL — 6 thành viên BẮT BUỘC điền sau khi tự mở production + DOCX cũ."),
        ("Cam", "PM/BA chốt requirement, Jira, ưu tiên, effort — sau khi đủ HITL."),
        ("Vàng", "Ô bắt buộc còn trống = tab chưa xong, không được báo 'đã UAT'."),
    ]
    ws["A4"] = "Màu"
    ws["B4"] = "Ý nghĩa — ai điền"
    style_header(ws, 4, ["a", "b"], BLUE)
    for i, (a, b) in enumerate(legend, 5):
        write_row(ws, i, [a, b], GREY if i % 2 else WHITE)
    rules = [
        "Môi trường UAT chính: User " + PROD_USER + " — Admin " + PROD_ADMIN,
        "Staging (cùng VM): " + STG_USER + " / " + STG_ADMIN + " — dùng khi Prod hotfix, ghi rõ môi trường.",
        "Local: user :4002, admin :4001 (localhost cho Google SSO), API :8787 — không thay Prod.",
        "Mỗi người: (1) đọc DOCX/BPMN cũ của tab mình (2) chạy UI production (3) điền HITL trên ĐÚNG tab tính năng (4) reviewer chấm lại.",
        "Cấm: điền hộ tab người khác; copypaste 'OK' vào mọi dòng; đề xuất AI không gắn function; viết ADR.md dài thay vì cột Why trên Jira/MR.",
        "Angular = SPA. API = Node. Đề xuất frontend phải Signals + HTTP trong service. Admin change-password hiện stub (KAN-9).",
        "Gap đã biết: KAN-7 lazy routes, KAN-8 empty/error cart-wishlist-checkout, KAN-9 đổi MK admin, U-29 notifications API không có UI chuông.",
        "Sau workshop: PM điền Requirement_ID + Jira_KAN. Developer lấy git log --grep=KAN-n, không lấy 'thứ Ba nói'.",
        "6 thành viên: Gia, Khải, Quỳnh, Uyên, Ninh, Hân — phân công tab 02_PhanCong_6TV. Reviewer = người kế bên (HITL kép).",
        "Đặt tên tab: U_* storefront, A_* admin. Master_Catalog là mục lục. Điền chi tiết trên tab con.",
    ]
    ws["A10"] = "Quy tắc"
    ws["A10"].font = SECTION_FONT
    for i, t in enumerate(rules, 11):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=8)
        write_row(ws, i, [t], YELLOW_CELL if i % 2 == 0 else WHITE)
    autosize(ws, [28, 90, 20, 20, 20, 20, 20, 20])
    ws.auto_filter.ref = "A4:B8"

    # ----- 01 ý nghĩa cột -----
    ws = wb.create_sheet("01_YNghiaCot")
    banner(
        ws,
        "Từ điển cột — ý định từng cột (đọc trước khi điền tab tính năng)",
        "Cột xanh = hệ thống đã điền. Cột xanh lá = human-in-the-loop. Cam = BA/PM. Mọi tab U_* và A_* dùng cùng schema UAT.",
    )
    headers = ["STT", "Ten_cot", "Ai_dien", "Bat_buoc", "Y_dinh_BA_PM", "Vi_du", "O_nao_duoc_de_trong"]
    style_header(ws, 4, headers, BLUE)
    for i, h in enumerate(headers, 1):
        ws.cell(4, i, h)
        ws.cell(4, i).fill = BLUE
        ws.cell(4, i).font = NAVY_FONT
    for i, (name, who, intent, ex) in enumerate(UAT_COLUMNS, 1):
        must = "Có" if who == "HITL" else ("Khi chốt" if who == "PM" else "Giữ đúng")
        empty_ok = "Không (trừ N/A có lý do)" if who == "HITL" else ("Được đến workshop" if who == "PM" else "Không xóa mã")
        fill = GREEN_CELL if who == "HITL" else (ORANGE_CELL if who == "PM" else BLUE_CELL)
        write_row(ws, 4 + i, [i, name, who, must, intent, ex, empty_ok], fill)
        ws.cell(4 + i, 2).comment = Comment(intent, "BA")
    r0 = 5 + len(UAT_COLUMNS)
    ws.cell(r0, 1, "Bảng HITL tiêu chí tab (cuối mỗi tab tính năng)")
    ws.cell(r0, 1).font = SECTION_FONT
    style_header(ws, r0 + 1, ["Ma", "Cau_hoi", "Y_dinh"], GREEN)
    ws.cell(r0 + 1, 1, "Ma")
    ws.cell(r0 + 1, 2, "Cau_hoi_HITL")
    ws.cell(r0 + 1, 3, "Y_dinh")
    for j, (ma, q, intent) in enumerate(HITL_CRITERIA, 1):
        write_row(ws, r0 + 1 + j, [ma, q, intent], GREEN_CELL)
    autosize(ws, [8, 22, 12, 16, 70, 40, 28])
    ws.auto_filter.ref = f"A4:G{4 + len(UAT_COLUMNS)}"
    ws.freeze_panes = "A5"

    # ----- 02 phân công -----
    ws = wb.create_sheet("02_PhanCong_6TV")
    banner(
        ws,
        "Phân công đồng đều 6 thành viên — mỗi người verify DOCX cũ + production",
        "Primary chạy UAT + điền HITL. Reviewer_HITL không được là cùng một người. Đếm tab phải đủ 43 function.",
    )
    h = ["Thanh_vien", "Vai_tro_goiy", "So_tab", "Tab_phu_trach", "Reviewer_HITL", "Khu_vuc", "DOCX_cu_can_mo", "Production_URL_chinh", "Tien_do_%", "Ghi_chu"]
    style_header(ws, 4, h, BLUE)
    for i, x in enumerate(h, 1):
        ws.cell(4, i, x).fill = BLUE
        ws.cell(4, i).font = NAVY_FONT
    clusters = {
        "Gia": ("Auth identity", "DOCX đăng nhập/đăng ký/SSO/đổi MK", f"{PROD_USER}/auth/signin + {PROD_ADMIN}/login"),
        "Khải": ("Catalog + giá + wishlist", "DOCX SP/BST/giá", f"{PROD_USER}/products + {PROD_ADMIN}/products"),
        "Quỳnh": ("Giỏ + checkout + KM", "DOCX thanh toán/voucher/freeship", f"{PROD_USER}/cart + {PROD_ADMIN}/promotions"),
        "Uyên": ("Đơn + đổi trả + onboard admin", "DOCX order-to-return, cấp quyền admin", f"{PROD_USER}/account/orders + {PROD_ADMIN}/orders"),
        "Ninh": ("AI + blog + dashboard + notify", "DOCX chatbot/quiz/dashboard", f"{PROD_USER}/chatbot + {PROD_ADMIN}/dashboard"),
        "Hân": ("Content + review + account + audit", "DOCX CSKH/review/policy", f"{PROD_USER}/policies + {PROD_ADMIN}/reviews"),
    }
    r = 5
    for m in MEMBERS:
        tabs = [f["sheet"] for f in FEATURES if f["owner"] == m]
        cluster, docx, url = clusters[m]
        write_row(
            ws,
            r,
            [m, cluster, len(tabs), ", ".join(tabs), SECONDARY[m], cluster, docx, url, 0, "Tự chạy, không chia nhỏ cho người ngoài 6TV"],
            GREEN_CELL,
        )
        r += 1
    ws.cell(r + 1, 1, "Tổng tab function")
    ws.cell(r + 1, 3, f"=COUNTA(C5:C10)")  # wrong - better hardcoded
    ws.cell(r + 1, 3, len(FEATURES))
    ws.cell(r + 2, 1, "Công thức tiến độ: mỗi người cập nhật cột Tien_do_% khi xong HITL H12. Dashboard 10_TienDo đọc sheet này.")
    ws.merge_cells(start_row=r + 2, start_column=1, end_row=r + 2, end_column=8)
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', "A5:A10")
    autosize(ws, [14, 28, 10, 70, 14, 28, 40, 55, 12, 40])

    # ----- 03 catalog -----
    ws = wb.create_sheet("03_Master_Catalog")
    banner(
        ws,
        "Mục lục toàn bộ tính năng hiện trạng (1 hàng = 1 tab UAT)",
        "Hyperlink cột Tab mở đúng sheet. Cột vàng/HITL điền trên tab con, không điền hộ ở đây trừ tiến độ.",
    )
    ch = [
        "Ma", "Tab", "Ung_dung", "Ten_tinh_nang", "Primary", "Reviewer", "Route_Angular",
        "URL_Production", "API_chinh", "Userflow_BPMN_hien_tai", "UI_hien_tai", "Loi_gap_da_biet",
        "So_TC", "Trang_thai_UAT", "Diem_HITL_1_5", "Co_DOCX",
    ]
    style_header(ws, 4, ch, BLUE)
    for i, x in enumerate(ch, 1):
        c = ws.cell(4, i, x)
        c.fill = BLUE
        c.font = NAVY_FONT
    for i, f in enumerate(FEATURES, 5):
        ncase = len(default_cases(f))
        vals = [
            f["ma"], f["sheet"], f["app"], f["ten"], f["owner"], f["reviewer"], f["route"],
            f["url"], f["api"], f["bpmn"], f["ui_now"], f["known"], ncase,
            "Chưa test", "", "Chưa xác nhận",
        ]
        fill = BLUE_CELL if f["app"] == "User" else ORANGE_CELL
        write_row(ws, i, vals, fill)
        ws.cell(i, 2).hyperlink = f"#'{f['sheet']}'!A1"
        ws.cell(i, 2).font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        ws.cell(i, 14).fill = YELLOW_CELL
        ws.cell(i, 15).fill = GREEN_CELL
        ws.cell(i, 16).fill = GREEN_CELL
    last = 4 + len(FEATURES)
    add_dv(ws, '"Chưa test,Đang test,Đạt,Không đạt,Blocked,Cần review HITL"', f"N5:N{last}")
    add_dv(ws, '"1,2,3,4,5"', f"O5:O{last}")
    add_dv(ws, '"Có — đã đọc,Không có DOCX,Chưa mở"', f"P5:P{last}")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', f"E5:F{last}")
    ws.auto_filter.ref = f"A4:P{last}"
    ws.freeze_panes = "A5"
    autosize(ws, [10, 20, 12, 28, 12, 12, 28, 42, 55, 40, 40, 45, 10, 18, 14, 18])

    # ----- 04 UAT tổng hợp (index of all cases) -----
    ws = wb.create_sheet("04_UAT_TongHop")
    banner(
        ws,
        "Toàn bộ test case — điền KẾT QUẢ trên tab tính năng; sheet này để PM lọc/P0",
        "Ma_TC trùng tab con. Sau khi tester xong, PM copy Pass_Fail vào đây hoặc lọc theo Primary.",
    )
    th = [c[0] for c in UAT_COLUMNS]
    style_header(ws, 4, th, BLUE)
    for i, name in enumerate(th, 1):
        cell = ws.cell(4, i, name)
        cell.fill = GREEN if UAT_COLUMNS[i - 1][1] == "HITL" else (ORANGE if UAT_COLUMNS[i - 1][1] == "PM" else BLUE)
        cell.font = NAVY_FONT
        cell.comment = Comment(UAT_COLUMNS[i - 1][2], "BA")
    row = 5
    all_cases = []
    for f in FEATURES:
        for idx, (nhom, con, steps, expected) in enumerate(default_cases(f), 1):
            ma_tc = f"{f['ma']}-TC{idx:02d}"
            all_cases.append((f, ma_tc, nhom, con, steps, expected, idx))
            vals = [""] * len(th)
            m = {c[0]: i for i, c in enumerate(UAT_COLUMNS)}
            def setc(k, v):
                vals[m[k]] = v
            setc("Ma_TC", ma_tc)
            setc("Nhom_TC", nhom)
            setc("Tinh_nang_con", con)
            setc("Precondition", f["workflow"] if nhom == "Happy" else "Xem Buoc_thao_tac")
            setc("Buoc_thao_tac", f"[{f['app']}] {f['url']} | {steps}")
            setc("Du_lieu_test", "")
            setc("Ket_qua_mong_doi", expected)
            setc("Tester", f["owner"])
            setc("Reviewer_HITL", f["reviewer"])
            setc("Trang_thai", "Chưa test")
            setc("Moi_truong", "Prod user" if f["app"] == "User" else "Prod admin")
            fill_row = YELLOW_CELL
            write_row(ws, row, vals, fill_row)
            # color SYS vs HITL cells
            for ci, col in enumerate(UAT_COLUMNS, 1):
                if col[1] == "SYS":
                    ws.cell(row, ci).fill = BLUE_CELL
                elif col[1] == "HITL":
                    ws.cell(row, ci).fill = YELLOW_CELL
                else:
                    ws.cell(row, ci).fill = ORANGE_CELL
            row += 1
    last_uat = row - 1
    add_dv(ws, '"Happy,Negative,Edge,Empty,Permission,Mobile,DOCX,AI"', f"B5:B{last_uat}")
    add_dv(ws, '"Pass,Fail,Blocked,N/A"', f"I5:I{last_uat}")
    add_dv(ws, '"Blocker,Major,Minor,Cosmetic"', f"K5:K{last_uat}")
    add_dv(ws, '"Prod user,Prod admin,Staging user,Staging admin,Local"', f"M5:M{last_uat}")
    add_dv(ws, '"Desktop 1280,Mobile 390,Tablet"', f"N5:N{last_uat}")
    add_dv(ws, '"Khớp,Lệch,Thiếu trên UI,Thiếu trên DOCX,Không có DOCX"', f"O5:O{last_uat}")
    add_dv(ws, '"Không,Thấp,Trung bình,Cao,Phá vỡ luồng"', f"T5:T{last_uat}")
    add_dv(ws, '"P0,P1,P2,P3"', f"AC5:AC{last_uat}")
    add_dv(ws, '"S,M,L,XL"', f"AD5:AD{last_uat}")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', f"AE5:AF{last_uat}")
    add_dv(ws, '"Chưa test,Đang test,Đạt,Không đạt,Blocked,Cần review HITL"', f"AH5:AH{last_uat}")
    ws.auto_filter.ref = f"A4:{get_column_letter(len(th))}{last_uat}"
    ws.freeze_panes = "A5"
    autosize(ws, [14, 12, 22, 28, 40, 22, 36, 28, 12, 28, 12, 18, 14, 14, 18, 28, 24, 28, 22, 16, 22, 22, 22, 24, 24, 24, 24, 16, 10, 10, 12, 12, 12, 16, 12])

    # ----- 05 bugs -----
    ws = wb.create_sheet("05_Loi_HienTrang")
    banner(ws, "Sổ lỗi hiện trạng (seed từ COMPLIANCE + phát hiện UAT)", "Tester thêm dòng khi Fail. Không xóa seed. PM map Jira sau.")
    bh = ["Ma_loi", "Nguon", "Tab", "Mo_ta", "Buoc_tai_hien", "Muc_do", "Tac_dong_user_admin", "Workaround", "De_xuat", "Jira", "Tester", "Trang_thai"]
    style_header(ws, 4, bh, BLUE)
    for i, x in enumerate(bh, 1):
        ws.cell(4, i, x).fill = BLUE
        ws.cell(4, i).font = NAVY_FONT
    seeds = [
        ("BUG-001", "COMPLIANCE", "U_GioHang/U_Wishlist/U_CK_*", "Empty/error UI chưa đủ", "Giỏ/wishlist/checkout rỗng hoặc API fail", "Major", "User không biết bước tiếp", "Refresh", "KAN-8 empty/error pattern", "KAN-8", "Quỳnh", "Đã biết"),
        ("BUG-002", "COMPLIANCE", "A_DoiMatKhau", "Đổi MK chỉ validate local, không gọi API", "Submit form ≥12 ký tự khớp", "Major", "Admin tưởng đã đổi MK", "Dùng forgot trên login", "Gọi API, bỏ stub", "KAN-9", "Gia", "Đã biết"),
        ("BUG-003", "COMPLIANCE", "Toàn SPA", "Chưa lazy loadComponent", "Mở network lúc boot", "Minor", "Bundle lớn, chậm mobile", "Không", "KAN-7", "KAN-7", "Ninh", "Đã biết"),
        ("BUG-004", "CODE", "U_ThongBao", "API notifications có, header không gắn chuông", "Login member, nhìn header", "Major", "Trượt thông báo đơn/KM", "Vào đơn thủ công", "Gắn bell + page", "", "Ninh", "Đã biết"),
        ("BUG-005", "CODE", "U_Home", "Search chỉ khi Enter, không suggest", "Gõ search không Enter", "Minor", "Lệch app fashion production", "Enter", "Typeahead + empty", "", "Khải", "Đã biết"),
        ("BUG-006", "CODE", "U_AI_Chatbot", "Policy fallback cứng nếu DB tool fail", "Hỏi đổi trả khi tool lỗi", "Major", "Lệch policy thật", "Mở /policies", "Fail closed + CSKH", "", "Ninh", "Đã biết"),
        ("BUG-007", "CODE", "A_Login", "Google callback phải đúng origin admin", "SSO từ domain sai", "Blocker (nếu sai env)", "Không vào admin", "Login password", "Document origin", "", "Gia", "Đã biết"),
    ]
    for i, s in enumerate(seeds, 5):
        write_row(ws, i, list(s), RED_CELL)
    for i in range(12, 40):
        write_row(ws, i, [f"BUG-{i-4:03d}"] + [""] * 11, YELLOW_CELL)
        ws.cell(i, 1).fill = YELLOW_CELL
    add_dv(ws, '"Blocker,Major,Minor,Cosmetic"', "F5:F40")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', "K5:K40")
    add_dv(ws, '"Đã biết,Mới phát hiện,Won\'t fix,Đã Jira,Đã sửa staging"', "L5:L40")
    ws.auto_filter.ref = "A4:L40"
    autosize(ws, [12, 14, 22, 40, 36, 12, 28, 22, 28, 12, 12, 18])

    # ----- 06 UI -----
    ws = wb.create_sheet("06_CaiTien_UIUX")
    banner(ws, "Cải tiến Wireframe / Mockup / Prototype / Frontend", "1 hàng = 1 đề xuất UI. Bắt buộc lý do + tác động + tab nguồn.")
    uh = ["Ma_UI", "Tab", "Lop", "Hien_trang", "De_xuat", "Ly_do", "Tac_dong", "Do_uu_tien_goi_y", "Primary", "Link_Figma", "Chot_PM"]
    style_header(ws, 4, uh, ORANGE)
    for i, x in enumerate(uh, 1):
        ws.cell(4, i, x).fill = ORANGE
        ws.cell(4, i).font = NAVY_FONT
    ui_seed = [
        ("UI-01", "U_Home", "Wireframe", "Search không gợi ý; không chuông", "Hàng search typeahead + bell", "Giống store production", "Trung bình", "P1", "Khải", "", ""),
        ("UI-02", "U_GioHang", "Frontend", "Empty yếu (KAN-8)", "Empty illustration + CTA + error retry", "Bỏ cuộc checkout", "Cao", "P0", "Quỳnh", "", ""),
        ("UI-03", "U_CK_VanChuyen", "Prototype", "Guest OTP vs member 1 bước dễ rối", "Stepper 3 bước rõ guest/member", "Giống checkout Shopee/fashion", "Cao", "P0", "Quỳnh", "", ""),
        ("UI-04", "U_AI_Chatbot", "Mockup", "Trang + widget trùng", "1 surface; PDP cards trong chat", "Stylist production", "Trung bình", "P1", "Ninh", "", ""),
        ("UI-05", "A_DoiTra_CSKH", "Wireframe", "4 zone trên 1 page nặng", "Inbox + return queue tách, unread", "CSKH ca trực", "Cao", "P1", "Uyên", "", ""),
        ("UI-06", "A_Dashboard", "Frontend", "Reload full page 'Làm mới'", "Refresh signal, không location.reload", "Mất filter", "Thấp", "P2", "Ninh", "", ""),
        ("UI-07", "U_ThongBao", "Prototype", "Không có UI", "Bell + drawer + deep link đơn", "API đã có", "Cao", "P0", "Ninh", "", ""),
        ("UI-08", "A_DoiMatKhau", "Frontend", "Stub", "Form gọi API + strength meter", "KAN-9", "Cao", "P0", "Gia", "", ""),
    ]
    for i, s in enumerate(ui_seed, 5):
        write_row(ws, i, list(s), ORANGE_CELL)
    for i in range(13, 50):
        write_row(ws, i, [f"UI-{i-4:02d}"] + [""] * 10, YELLOW_CELL)
    add_dv(ws, '"Wireframe,Mockup,Prototype,Frontend,Token/a11y"', "C5:C50")
    add_dv(ws, '"Không,Thấp,Trung bình,Cao,Phá vỡ luồng"', "G5:G50")
    add_dv(ws, '"P0,P1,P2,P3"', "H5:H50")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', "I5:I50")
    add_dv(ws, '"Nháp,HITL xong,PM chốt,Làm design,Dev"', "K5:K50")
    ws.auto_filter.ref = "A4:K50"
    autosize(ws, [10, 18, 14, 36, 40, 28, 16, 14, 12, 22, 14])

    # ----- 07 BPMN -----
    ws = wb.create_sheet("07_BPMN_Userflow")
    banner(ws, "Cải tiến BPMN + userflow lối vào tính năng", "Lane bắt buộc: User | Admin | Hệ thống/API | Payment | CSKH. Ghi cổng hiện tại vs cổng đề xuất.")
    ph = ["Ma_QT", "Tab", "Ten_quy_trinh", "BPMN_hien_tai", "Lo_hong", "BPMN_de_xuat", "Userflow_vao_tinh_nang_hien_tai", "Userflow_giong_production", "Exception", "Tac_dong", "Primary"]
    style_header(ws, 4, ph, BLUE)
    for i, x in enumerate(ph, 1):
        ws.cell(4, i, x).fill = BLUE
        ws.cell(4, i).font = NAVY_FONT
    proc = [
        ("QT-01", "U_CK_*", "Guest/Member checkout", "Cart→ship→(OTP|pay)→confirm", "F5 mất confirm; PTTT pending", "Thêm wait payment + resume bằng tracking", "CTA giỏ", "Mini-cart + progress + email resume", "OTP hết hạn, cổng fail", "Cao", "Quỳnh"),
        ("QT-02", "U_TK_DoiTra+A_DoiTra", "Đổi trả 48h", "Delivered→user request→admin refund/exchange/reject", "Evidence, vận chuyển hoàn, cổng hoàn tiền", "Thêm QC inbound + refund gateway task", "Từ order detail", "Deep link từ notify + chatbot", "Ngoài 48h, sai tem", "Cao", "Uyên"),
        ("QT-03", "U_Auth+A_*", "Cấp quyền admin", "Register→welcome→role request→duyệt", "User tưởng đã là admin", "Email khi duyệt + force change-password API", "URL /register", "Invite link hết hạn", "Tự gõ /dashboard", "Trung bình", "Gia"),
        ("QT-04", "U_AI_*", "Quiz→gợi ý→chat", "Modal invite, 2 bề mặt chat", "Mất context quiz trong chat", "Một profile stylist xuyên suốt", "Nav Gợi ý AI / Chatbot", "Entry từ PDP 'hỏi stylist'", "Guest migrate quiz", "Trung bình", "Ninh"),
        ("QT-05", "A_DonHang", "Status đơn", "change-status/cancel/resolve pay", "Nhảy status", "Máy trạng thái + audit bắt buộc", "Dashboard pending", "Hàng đợi ca + SLA", "Payment error", "Cao", "Uyên"),
        ("QT-06", "U_TK_DanhGia+A_DanhGia", "Review moderation", "POST review→approve/hide", "Không AI pre-screen", "AI flag + human approve", "My reviews", "Sau delivered notify", "Spam", "Thấp", "Hân"),
    ]
    for i, s in enumerate(proc, 5):
        write_row(ws, i, list(s), BLUE_CELL)
    for i in range(11, 40):
        write_row(ws, i, [f"QT-{i-4:02d}"] + [""] * 10, YELLOW_CELL)
    add_dv(ws, '"Không,Thấp,Trung bình,Cao,Phá vỡ luồng"', "J5:J40")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', "K5:K40")
    ws.auto_filter.ref = "A4:K40"
    autosize(ws, [10, 18, 24, 36, 32, 40, 28, 36, 24, 14, 12])

    # ----- 08 AI -----
    ws = wb.create_sheet("08_AI_TinhNangMoi")
    banner(ws, "Tính năng AI mới — gắn từng function (không AI chung)", "Mỗi đề xuất: input, output, human approve, rủi ro, tab. HITL chấm có làm hay không.")
    ah = ["Ma_AI", "Tab_gan", "Ten", "Input", "Output", "Human_in_the_loop", "Rui_ro", "Hien_da_co", "De_xuat", "Ly_do", "Tac_dong", "Primary", "Prio"]
    style_header(ws, 4, ah, GREEN)
    for i, x in enumerate(ah, 1):
        ws.cell(4, i, x).fill = GREEN
        ws.cell(4, i).font = NAVY_FONT
    ai_rows = [
        ("AI-01", "U_AI_Chatbot", "Stylist tool-use siết policy", "Câu user + session", "Trả lời + product cards", "CSKH takeover A-11", "Bịa policy", "LLM + tools", "Bắt buộc tool policy; handoff human", "Đã có prompt, cần UX handoff", "Trung bình", "Ninh", "P1"),
        ("AI-02", "U_AI_GoiY", "Outfit combo giải thích", "Quiz profile", "Combo + reason", "User sửa tag", "Gợi ý sai size", "recommendations/style-profile", "Size/in-stock filter + explain", "Production stylist", "Cao", "Ninh", "P1"),
        ("AI-03", "U_SP_DanhSach", "Search semantic + typeahead", "q", "SP + query rewrite", "Không (log)", "SP lệch", "q exact Enter", "Embedding search", "Search hiện yếu", "Cao", "Khải", "P1"),
        ("AI-04", "A_DanhGia", "Pre-moderation độc hại", "Text+ảnh review", "Flag", "Admin approve bắt buộc", "False positive", "Duyệt tay", "Model + queue", "Giảm leak PDP", "Thấp", "Hân", "P2"),
        ("AI-05", "A_DoiTra_CSKH", "Tóm tắt ticket/chat", "Thread", "Tóm tắt + next action", "Agent sửa trước gửi", "Sai hoàn tiền", "Chat list", "Summarize + draft reply", "Ca trực", "Cao", "Uyên", "P1"),
        ("AI-06", "A_Dashboard", "Anomaly đơn/thanh toán", "Timeseries", "Cảnh báo", "Ops xác nhận", "Cảnh báo giả", "KPI thô", "Threshold + LLM giải thích", "Ops", "Trung bình", "Ninh", "P2"),
        ("AI-07", "U_CK_VanChuyen", "Address autocomplete + risk COD", "Địa chỉ", "Chuẩn hóa + risk", "User xác nhận địa chỉ", "Giao fail", "Tỉnh/huyện/xã tay", "Suggest + COD risk", "Giống production TMĐT", "Cao", "Quỳnh", "P1"),
        ("AI-08", "A_SanPham", "Mô tả SP / tag từ ảnh", "Ảnh+tên", "Mô tả/tag nháp", "Merch duyệt", "Sai chất liệu", "Nhập tay", "Draft + approve", "Tốc độ catalog", "Trung bình", "Khải", "P2"),
        ("AI-09", "U_ThongBao", "Notify thông minh", "Event đơn", "Bell + copy", "Template PM", "Spam", "API chưa UI", "UI + digest", "API sẵn", "Cao", "Ninh", "P0"),
        ("AI-10", "U_Auth_DangNhap", "Risk login", "IP/device", "Step-up OTP", "Không", "Khóa nhầm", "OTP unverified", "Device signal", "Bảo mật", "Thấp", "Gia", "P2"),
    ]
    for i, s in enumerate(ai_rows, 5):
        write_row(ws, i, list(s), GREEN_CELL)
    for i in range(15, 40):
        write_row(ws, i, [f"AI-{i-4:02d}"] + [""] * 12, YELLOW_CELL)
    add_dv(ws, '"Không,Thấp,Trung bình,Cao,Phá vỡ luồng"', "K5:K40")
    add_dv(ws, '"Gia,Khải,Quỳnh,Uyên,Ninh,Hân"', "L5:L40")
    add_dv(ws, '"P0,P1,P2,P3"', "M5:M40")
    ws.auto_filter.ref = "A4:M40"
    autosize(ws, [10, 16, 28, 22, 24, 24, 18, 22, 36, 24, 14, 12, 10])

    # ----- 09 HITL scoreboard -----
    ws = wb.create_sheet("09_HITL_ChamDiem")
    banner(ws, "Human-in-the-loop — chấm từng tab (Primary điền, Reviewer xác nhận)", "Điểm 1–5. Cột vàng bắt buộc. Không trung bình hộ — reviewer được phép lệch và phải ghi lý do.")
    hh = ["Ma", "Tab", "Primary", "Reviewer", "H1_DOCX", "H2_UI", "H3_Wire", "H4_Mock", "H5_Proto", "H6_FE", "H7_BPMN", "H8_Flow", "H9_AI", "H10_RuiRo", "H11_Diem_Primary", "H11_Diem_Reviewer", "H12_Chot", "Nhan_xet_khep"]
    style_header(ws, 4, hh, GREEN)
    for i, x in enumerate(hh, 1):
        ws.cell(4, i, x).fill = GREEN
        ws.cell(4, i).font = NAVY_FONT
        if i >= 5:
            ws.cell(4, i).comment = Comment(HITL_CRITERIA[min(i - 5, len(HITL_CRITERIA) - 1)][1], "BA")
    for i, f in enumerate(FEATURES, 5):
        vals = [f["ma"], f["sheet"], f["owner"], f["reviewer"]] + [""] * 14
        write_row(ws, i, vals, YELLOW_CELL)
        for c in range(1, 5):
            ws.cell(i, c).fill = BLUE_CELL
    last_h = 4 + len(FEATURES)
    add_dv(ws, '"1,2,3,4,5"', f"E5:P{last_h}")
    add_dv(ws, '"Có,Chưa,Cần workshop"', f"Q5:Q{last_h}")
    ws.auto_filter.ref = f"A4:R{last_h}"
    ws.freeze_panes = "A5"
    autosize(ws, [10, 20, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12, 16, 16, 14, 40])

    # ----- 10 tiến độ -----
    ws = wb.create_sheet("10_TienDo")
    banner(ws, "Tiến độ khảo sát — cập nhật % trên 02_PhanCong và Pass_Fail trên tab", "Biểu đồ đếm tab theo owner. Công thức đếm Pass trên 04_UAT_TongHop.")
    ws["A4"] = "Thành viên"
    ws["B4"] = "Số tab"
    ws["C4"] = "Tiến độ % (từ phân công)"
    style_header(ws, 4, ["a", "b", "c"], BLUE)
    for i, m in enumerate(MEMBERS, 5):
        n = sum(1 for f in FEATURES if f["owner"] == m)
        ws.cell(i, 1, m)
        ws.cell(i, 2, n)
        ws.cell(i, 3, f"='02_PhanCong_6TV'!I{i}")
        write_row(ws, i, [m, n, f"='02_PhanCong_6TV'!I{i}"], GREEN_CELL)
    ws["A12"] = "Tổng test case"
    ws["B12"] = len(all_cases)
    ws["A13"] = "Tổng function tab"
    ws["B13"] = len(FEATURES)
    ws["A14"] = "Pass (cột I sheet 04, khi đã điền)"
    ws["B14"] = f'=COUNTIF(\'04_UAT_TongHop\'!I:I,"Pass")'
    ws["A15"] = "Fail"
    ws["B15"] = f'=COUNTIF(\'04_UAT_TongHop\'!I:I,"Fail")'
    ws["A16"] = "Chưa test"
    ws["B16"] = f'=COUNTIF(\'04_UAT_TongHop\'!AH:AH,"Chưa test")'
    chart = BarChart()
    chart.title = "Số tab UAT theo thành viên"
    chart.y_axis.title = "Số tab"
    chart.x_axis.title = "Thành viên"
    data = Reference(ws, min_col=2, min_row=4, max_row=10)
    cats = Reference(ws, min_col=1, min_row=5, max_row=10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    ws.add_chart(chart, "E4")
    autosize(ws, [36, 18, 28])

    # ----- feature sheets -----
    uat_names = [c[0] for c in UAT_COLUMNS]
    for f in FEATURES:
        w = wb.create_sheet(f["sheet"][:31])
        app_label = "STOREFRONT USER" if f["app"] == "User" else "ADMIN"
        banner(
            w,
            f"{f['ma']}  {f['ten']}  [{app_label}]  — Owner {f['owner']}  /  Reviewer {f['reviewer']}",
            f"URL: {f['url']}   |   Route: {f['route']}   |   Đọc 01_YNghiaCot trước. Ô vàng = phải điền. Không điền tab người khác.",
        )
        meta = [
            ("Ma", f["ma"]),
            ("Ung_dung", f["app"]),
            ("Primary / Reviewer HITL", f"{f['owner']} / {f['reviewer']}"),
            ("Route Angular", f["route"]),
            ("URL production", f["url"]),
            ("API / hợp đồng code", f["api"]),
            ("Userflow / BPMN hiện tại", f["bpmn"]),
            ("UI đang chạy", f["ui_now"]),
            ("Workflow chi tiết", f["workflow"]),
            ("Lỗi / gap đã biết từ code", f["known"]),
            ("DOCX cũ (HITL dán đường dẫn)", ""),
            ("Ngày test production", ""),
        ]
        w["A3"] = "Hiện trạng từ CODE (SYS) — chỉ sửa khi sai"
        w["A3"].font = SECTION_FONT
        for i, (k, v) in enumerate(meta):
            rr = 4 + i
            w.cell(rr, 1, k).fill = BLUE
            w.cell(rr, 1).font = NAVY_FONT
            w.cell(rr, 1).border = THIN
            w.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
            cell = w.cell(rr, 2, v)
            cell.fill = YELLOW_CELL if v == "" else BLUE_CELL
            cell.alignment = WRAP
            cell.border = THIN
            w.row_dimensions[rr].height = 32
        start = 4 + len(meta) + 2
        w.cell(start, 1, "PHẦN A — UAT từng bước trên production (mỗi hàng = 1 test case)")
        w.cell(start, 1).font = SECTION_FONT
        href = start + 1
        style_header(w, href, uat_names, BLUE)
        for i, col in enumerate(UAT_COLUMNS, 1):
            cell = w.cell(href, i, col[0])
            cell.fill = GREEN if col[1] == "HITL" else (ORANGE if col[1] == "PM" else BLUE)
            cell.font = NAVY_FONT
            cell.comment = Comment(f"[{col[1]}] {col[2]}", "BA")
        cases = default_cases(f)
        first_data = href + 1
        for idx, (nhom, con, steps, expected) in enumerate(cases, 1):
            r = href + idx
            vals = [""] * len(uat_names)
            m = {c[0]: i for i, c in enumerate(UAT_COLUMNS)}
            vals[m["Ma_TC"]] = f"{f['ma']}-TC{idx:02d}"
            vals[m["Nhom_TC"]] = nhom
            vals[m["Tinh_nang_con"]] = con
            vals[m["Precondition"]] = "Xem workflow hiện trạng phía trên"
            vals[m["Buoc_thao_tac"]] = steps
            vals[m["Ket_qua_mong_doi"]] = expected
            vals[m["Tester"]] = f["owner"]
            vals[m["Reviewer_HITL"]] = f["reviewer"]
            vals[m["Trang_thai"]] = "Chưa test"
            vals[m["Moi_truong"]] = "Prod user" if f["app"] == "User" else "Prod admin"
            write_row(w, r, vals)
            for ci, col in enumerate(UAT_COLUMNS, 1):
                if col[1] == "SYS":
                    w.cell(r, ci).fill = BLUE_CELL
                elif col[1] == "HITL":
                    w.cell(r, ci).fill = YELLOW_CELL
                else:
                    w.cell(r, ci).fill = ORANGE_CELL
        last_c = href + len(cases)
        # extra blank HITL rows
        for extra in range(1, 4):
            r = last_c + extra
            write_row(w, r, [f"{f['ma']}-TC{len(cases)+extra:02d}"] + [""] * (len(uat_names) - 1), YELLOW_CELL)
            w.cell(r, m["Tester"] + 1, f["owner"])
            w.cell(r, m["Reviewer_HITL"] + 1, f["reviewer"])
        end_a = last_c + 3
        rng = f"I{first_data}:I{end_a}"
        add_dv(w, '"Pass,Fail,Blocked,N/A"', rng)
        add_dv(w, '"Blocker,Major,Minor,Cosmetic"', f"K{first_data}:K{end_a}")
        add_dv(w, '"Prod user,Prod admin,Staging user,Staging admin,Local"', f"M{first_data}:M{end_a}")
        add_dv(w, '"Desktop 1280,Mobile 390,Tablet"', f"N{first_data}:N{end_a}")
        add_dv(w, '"Khớp,Lệch,Thiếu trên UI,Thiếu trên DOCX,Không có DOCX"', f"O{first_data}:O{end_a}")
        add_dv(w, '"Không,Thấp,Trung bình,Cao,Phá vỡ luồng"', f"T{first_data}:T{end_a}")
        add_dv(w, '"P0,P1,P2,P3"', f"AC{first_data}:AC{end_a}")
        add_dv(w, '"S,M,L,XL"', f"AD{first_data}:AD{end_a}")
        add_dv(w, '"Chưa test,Đang test,Đạt,Không đạt,Blocked,Cần review HITL"', f"AH{first_data}:AH{end_a}")
        w.auto_filter.ref = f"A{href}:{get_column_letter(len(uat_names))}{end_a}"
        w.freeze_panes = f"A{href+1}"

        hitl_row = end_a + 2
        w.cell(hitl_row, 1, "PHẦN B — HITL khảo sát cải tiến cấp TAB (Primary chấm, Reviewer xác nhận cột G)")
        w.cell(hitl_row, 1).font = SECTION_FONT
        hh2 = ["Ma_tieu_chi", "Cau_hoi", "Y_dinh_cot", "Diem_Primary_1_5", "Nhan_xet_Primary", "Diem_Reviewer_1_5", "Nhan_xet_Reviewer", "Bang_chung_URL"]
        hr = hitl_row + 1
        style_header(w, hr, hh2, GREEN)
        for i, x in enumerate(hh2, 1):
            w.cell(hr, i, x).fill = GREEN
            w.cell(hr, i).font = NAVY_FONT
        for j, (ma, q, intent) in enumerate(HITL_CRITERIA, 1):
            write_row(w, hr + j, [ma, q, intent, "", "", "", "", ""], GREEN_CELL)
            w.cell(hr + j, 4).fill = YELLOW_CELL
            w.cell(hr + j, 5).fill = YELLOW_CELL
            w.cell(hr + j, 6).fill = YELLOW_CELL
            w.cell(hr + j, 7).fill = YELLOW_CELL
        add_dv(w, '"1,2,3,4,5"', f"D{hr+1}:D{hr+len(HITL_CRITERIA)}")
        add_dv(w, '"1,2,3,4,5"', f"F{hr+1}:F{hr+len(HITL_CRITERIA)}")
        autosize(w, [16, 14, 22, 36, 44, 22, 40, 28, 12, 28, 12, 16, 14, 14, 18, 28, 24, 28, 22, 16, 22, 22, 22, 24, 24, 24, 24, 16, 10, 10, 12, 12, 12, 16, 12])
        w.sheet_view.showGridLines = False

    # print area / tab colors
    for f in FEATURES:
        sh = wb[f["sheet"][:31]]
        sh.sheet_properties.tabColor = "1F4E79" if f["app"] == "User" else "B9770E"

    wb["00_DocTruocKhiDien"].sheet_properties.tabColor = "196F3D"
    wb["01_YNghiaCot"].sheet_properties.tabColor = "196F3D"
    wb["02_PhanCong_6TV"].sheet_properties.tabColor = "196F3D"
    wb["03_Master_Catalog"].sheet_properties.tabColor = "1F4E79"
    wb["04_UAT_TongHop"].sheet_properties.tabColor = "1F4E79"
    wb["05_Loi_HienTrang"].sheet_properties.tabColor = "922B21"
    wb["06_CaiTien_UIUX"].sheet_properties.tabColor = "B9770E"
    wb["07_BPMN_Userflow"].sheet_properties.tabColor = "B9770E"
    wb["08_AI_TinhNangMoi"].sheet_properties.tabColor = "196F3D"
    wb["09_HITL_ChamDiem"].sheet_properties.tabColor = "196F3D"
    wb["10_TienDo"].sheet_properties.tabColor = "1F4E79"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    try:
        wb.save(DOWNLOADS)
    except OSError:
        pass
    return OUT, len(FEATURES), len(all_cases)


if __name__ == "__main__":
    path, nfeat, ncase = build()
    from collections import Counter

    counts = Counter(f["owner"] for f in FEATURES)
    print("Wrote", path)
    print("Features", nfeat, "cases", ncase)
    print("Owner counts", {k: counts[k] for k in ("Gia", "Khai", "Quynh", "Uyen", "Ninh", "Han")})
    print("Raw", [(f["owner"], 1) for f in FEATURES[:0]], "Gia", counts.get("Gia"), "Khai", counts.get("Khải"), "Quynh", counts.get("Quỳnh"), "Uyen", counts.get("Uyên"), "Ninh", counts.get("Ninh"), "Han", counts.get("Hân"))
