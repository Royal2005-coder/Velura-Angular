# -*- coding: utf-8 -*-
"""Bo sung tim kiem bang hinh anh (U6) vao docx muc 3.1.3.

Mo tep mau 'Quy trinh loc tim kiem san pham va them vao gio hang', GIU NGUYEN
moi noi dung cu, chi chen them:
  - 3 doan mo ta + Hinh 3.4b / 3.4c / 3.4d sau Hinh 3.4 (muc 3.1.3.1)
  - SHOP -06 .. SHOP -13 cuoi Bang 3.3 (muc 3.1.3.2)
  - Tinh huong 5 .. 11 sau Tinh huong 4 (muc 3.1.3.3)
Dinh dang lay bang cach nhan ban cac doan/hang da co cua mau.
"""
import copy
import glob
import os
import sys

from docx import Document
from docx.shared import Emu

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "build"))
from build_docx import normalise_measures, apply_pPr  # noqa: E402

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
HERE = os.path.dirname(os.path.abspath(__file__))
BA_DIR = os.path.abspath(os.path.join(HERE, ".."))
OUT = os.path.join(BA_DIR, "Velura-3.1.3-Tim-kiem-bang-hinh-anh.docx")

MO_TA = [
    ("Bên cạnh tìm kiếm bằng từ khóa và danh mục, Velura bổ sung cách tìm kiếm bằng hình ảnh theo kiểu "
     "Google Lens, giúp Guest và Member tìm được sản phẩm tương tự khi đã có sẵn một bức ảnh trang phục "
     "nhưng chưa biết gọi tên hay dùng từ khóa nào. Tại ô tìm kiếm trên thanh điều hướng, người dùng nhấn "
     "biểu tượng camera để mở cửa sổ chọn ảnh, có thể tải ảnh từ thiết bị, chụp trực tiếp, kéo thả hoặc "
     "dán ảnh. Hệ thống kiểm tra định dạng (JPG, PNG, WEBP) và dung lượng tối đa 5 MB trước khi xử lý; ảnh "
     "không hợp lệ bị từ chối ngay tại cửa sổ chọn ảnh kèm lý do cụ thể và không được gửi sang dịch vụ AI. "
     "Với ảnh hợp lệ, hệ thống hiển thị ảnh xem trước cùng khung chọn vùng (mặc định là toàn bộ ảnh) để "
     "người dùng khoanh đúng món đồ cần tìm khi ảnh có nhiều món.",
     "timkiem-anh-3-4b.png",
     "Hình 3.4b: BPMN quy trình chọn ảnh và kiểm tra ảnh trong tìm kiếm bằng hình ảnh"),
    ("Sau khi người dùng nhấn Tìm sản phẩm, hệ thống kiểm tra giới hạn tần suất của phiên rồi gửi ảnh tới "
     "dịch vụ AI để nhận diện trang phục và trả về các thuộc tính có cấu trúc gồm loại sản phẩm, màu sắc, "
     "phom dáng, chất liệu và phong cách. Nếu dịch vụ lỗi hoặc không phản hồi trong 15 giây (sau tối đa 2 "
     "lần thử lại), hệ thống báo lỗi tạm thời và cho phép thử lại; nếu ảnh không phải trang phục, hệ thống "
     "thông báo không nhận diện được và gợi ý chụp lại một món đồ rõ nét hoặc chuyển sang tìm bằng từ khóa. "
     "Với ảnh hợp lệ, các thuộc tính được ghép thành một mô tả văn bản, chuyển thành vector embedding rồi "
     "so khớp với vector của các sản phẩm đang bán theo độ tương đồng cosine (ngưỡng mặc định 0,45, tối đa "
     "20 sản phẩm). Ảnh chỉ được dùng để xử lý trong lượt tìm kiếm hiện tại và không được lưu lại sau khi "
     "có kết quả. Nếu không có sản phẩm nào đủ tương đồng, hệ thống hiển thị thông báo kèm khối "
     "“Sản phẩm nổi bật”.",
     "timkiem-anh-3-4c.png",
     "Hình 3.4c: BPMN quy trình phân tích ảnh và so khớp sản phẩm"),
    ("Để tối ưu với Style Quiz đã có trên web, danh sách tương đồng tiếp tục đi qua bước kiểm tra Style "
     "Profile như block “Gợi ý dành cho bạn” ở trang chủ. Nếu người dùng đã có Style Profile (Member lấy "
     "từ cơ sở dữ liệu, Guest lấy từ phiên), hệ thống chấm điểm lại kết quả theo dáng người, phong cách, "
     "tone da, dịp sử dụng và ngân sách, đồng thời gắn nhãn “Hợp dáng người của bạn”. Nếu chưa có, hệ thống "
     "giữ thứ tự theo độ tương đồng, khóa bộ lọc Dáng người theo quy tắc SHOP -02 và hiển thị banner mời "
     "làm Style Quiz. Trang kết quả gồm ảnh gốc thu nhỏ, các chip thuộc tính cho phép sửa hoặc xóa khi AI "
     "nhận diện sai, ô thêm từ khóa, các bộ lọc danh mục, giá, màu, size sẵn có và lưới sản phẩm. Khi người "
     "dùng tinh chỉnh, hệ thống tìm lại theo thuộc tính đã chỉnh mà không cần tải lại ảnh. Khi người dùng "
     "chọn một sản phẩm, các bước xem chi tiết, gợi ý size, chọn size, màu và Thêm vào giỏ hàng tiếp tục "
     "như mô tả ở Hình 3.4.",
     "timkiem-anh-3-4d.png",
     "Hình 3.4d: BPMN quy trình tối ưu theo Style Quiz, tinh chỉnh kết quả và thêm vào giỏ hàng"),
]

QUY_TAC = [
    ("SHOP -06", "Điểm vào tìm kiếm bằng hình ảnh",
     "Biểu tượng camera hiển thị trong ô tìm kiếm cho cả Guest và Member, không yêu cầu đăng nhập.",
     "Mở cửa sổ chọn ảnh với các cách tải lên, chụp ảnh, kéo thả hoặc dán; nếu quyền camera bị từ chối thì chuyển về tải ảnh lên."),
    ("SHOP -07", "Kiểm tra ảnh đầu vào",
     "Chỉ chấp nhận ảnh JPG, PNG, WEBP có dung lượng tối đa 5 MB.",
     "Từ chối ảnh sai định dạng hoặc quá nặng ngay tại cửa sổ chọn ảnh kèm lý do cụ thể; không gửi ảnh tới dịch vụ AI."),
    ("SHOP -08", "Chọn vùng tìm kiếm",
     "Vùng tìm mặc định là toàn bộ ảnh; người dùng có thể khoanh một món đồ khi ảnh có nhiều món.",
     "Cắt ảnh theo vùng đã chọn trước khi gửi phân tích; mỗi lượt tìm chỉ phân tích một món đồ."),
    ("SHOP -09", "Nhận diện trang phục",
     "Ảnh được phân tích để lấy thuộc tính gồm loại sản phẩm, màu, phom, chất liệu, phong cách. Ảnh không phải trang phục không được dùng để tìm kiếm.",
     "Gọi dịch vụ AI với thời gian chờ tối đa 15 giây, thử lại tối đa 2 lần; trả thuộc tính có cấu trúc, hoặc thông báo không nhận diện được / lỗi dịch vụ."),
    ("SHOP -10", "So khớp sản phẩm tương đồng",
     "Chỉ so khớp với sản phẩm đang bán và xếp theo độ tương đồng giảm dần.",
     "Tạo vector embedding từ mô tả thuộc tính, truy vấn với ngưỡng tương đồng cấu hình (mặc định 0,45), lấy tối đa 20 sản phẩm; nếu không có kết quả thì hiển thị khối Sản phẩm nổi bật."),
    ("SHOP -11", "Tối ưu theo Style Quiz",
     "Kết quả tìm bằng ảnh được cá nhân hóa theo Style Profile đang có hiệu lực (phiên của Guest hoặc CSDL của Member), cùng nguồn dữ liệu với SHOP -01; bộ lọc Dáng người vẫn tuân theo SHOP -02.",
     "Có Style Profile: chấm điểm lại theo dáng người, phong cách, tone da, dịp, ngân sách và gắn nhãn “Hợp dáng”. Chưa có: giữ thứ tự tương đồng, khóa lọc Dáng người và hiển thị banner mời làm Style Quiz."),
    ("SHOP -12", "Tinh chỉnh kết quả",
     "Người dùng được sửa hoặc xóa thuộc tính do AI nhận diện, thêm từ khóa và dùng các bộ lọc danh mục, giá, màu, size hiện có.",
     "Tìm lại theo thuộc tính và bộ lọc đã chỉnh mà không cần tải lại ảnh."),
    ("SHOP -13", "Bảo vệ quyền riêng tư và chi phí AI",
     "Ảnh chỉ dùng để tìm kiếm trong lượt hiện tại, không lưu lâu dài; số lượt tìm bằng ảnh được giới hạn theo phiên.",
     "Xóa ảnh sau khi trả kết quả, chỉ ghi nhật ký lượt tìm (không ghi nội dung ảnh); trả thông báo kèm thời gian chờ khi vượt giới hạn."),
]

TINH_HUONG = [
    ("Tình huống 5: Ảnh sai định dạng hoặc vượt quá dung lượng cho phép",
     "Mô tả: Người dùng chọn tệp không phải JPG, PNG, WEBP hoặc có dung lượng lớn hơn 5 MB trong cửa sổ chọn ảnh.",
     "Cách xử lý: Hệ thống không gửi ảnh đi xử lý, hiển thị thông báo ngay tại cửa sổ chọn ảnh: “Ảnh không hợp lệ. Vui lòng chọn ảnh JPG, PNG hoặc WEBP có dung lượng tối đa 5 MB.” và giữ người dùng ở lại cửa sổ để chọn ảnh khác."),
    ("Tình huống 6: Người dùng không cấp quyền camera",
     "Mô tả: Người dùng chọn chụp ảnh nhưng trình duyệt hoặc thiết bị từ chối quyền truy cập camera.",
     "Cách xử lý: Hệ thống hiển thị thông báo: “Không thể mở camera. Bạn có thể tải ảnh lên từ thiết bị.” và tự chuyển về chế độ tải ảnh lên, kéo thả hoặc dán ảnh."),
    ("Tình huống 7: Ảnh không phải trang phục hoặc không nhận diện được",
     "Mô tả: Ảnh là phong cảnh, thực phẩm, quá mờ hoặc có nhiều món đồ khiến dịch vụ AI không xác định được trang phục cần tìm.",
     "Cách xử lý: Hệ thống không chạy bước so khớp, hiển thị thông báo: “Chưa nhận diện được trang phục trong ảnh. Hãy thử ảnh rõ nét hơn, chỉ có một món đồ.” kèm hai lựa chọn “Thử ảnh khác” và “Tìm bằng từ khóa”."),
    ("Tình huống 8: Dịch vụ nhận diện ảnh bị lỗi hoặc quá thời gian chờ",
     "Mô tả: Dịch vụ AI không phản hồi trong 15 giây (sau tối đa 2 lần thử lại) hoặc trả lỗi.",
     "Cách xử lý: Hệ thống dừng lượt tìm, hiển thị thông báo: “Dịch vụ tìm kiếm bằng hình ảnh đang tạm thời gián đoạn. Vui lòng thử lại sau.” kèm nút “Thử lại”, đồng thời ghi nhận lỗi vào nhật ký hệ thống (Log) để giám sát."),
    ("Tình huống 9: Không có sản phẩm tương đồng với ảnh",
     "Mô tả: Trang phục được nhận diện nhưng không có sản phẩm đang bán nào đạt ngưỡng tương đồng.",
     "Cách xử lý: Hệ thống hiển thị thông báo: “Chưa tìm thấy sản phẩm giống với ảnh của bạn.” kèm khối “Sản phẩm nổi bật” và gợi ý chỉnh lại thuộc tính hoặc tìm bằng từ khóa."),
    ("Tình huống 10: Vượt giới hạn số lượt tìm bằng hình ảnh",
     "Mô tả: Một phiên gửi quá nhiều yêu cầu tìm bằng ảnh trong thời gian ngắn.",
     "Cách xử lý: Hệ thống từ chối lượt tìm mới và hiển thị thông báo: “Bạn đã tìm bằng hình ảnh quá nhiều lần. Vui lòng thử lại sau ít phút.” kèm thời gian chờ; tìm kiếm bằng từ khóa vẫn hoạt động bình thường."),
    ("Tình huống 11: Thành viên chưa có Hồ sơ phong cách dùng bộ lọc Dáng người trên kết quả tìm bằng ảnh",
     "Mô tả: Khách hàng đã đăng nhập nhưng chưa làm Style Quiz, muốn dùng bộ lọc Dáng người trên trang kết quả tìm bằng hình ảnh.",
     "Cách xử lý: Tương tự Tình huống 4, hệ thống không áp dụng bộ lọc và hiển thị popup mời thiết lập Hồ sơ phong cách; kết quả vẫn giữ nguyên thứ tự theo độ tương đồng với ảnh."),
]


def find_template():
    for path in glob.glob(os.path.join(BA_DIR, "Quy*.docx")):
        try:
            if Document(path).paragraphs[0].text.strip().startswith("3.1.3. Quy trình tìm kiếm"):
                return path
        except Exception:  # noqa: BLE001 - tep khong doc duoc thi bo qua
            continue
    raise SystemExit("khong tim thay tep mau 3.1.3")


def set_text(p_el, text):
    """Giu run dau (dinh dang), bo cac run con lai, dat noi dung moi."""
    runs = p_el.findall(W + "r")
    first = runs[0]
    for child in list(p_el):
        if child is not first and child.tag != W + "pPr":
            p_el.remove(child)       # run thua, hyperlink (link) cua chu thich cu
    for t in first.findall(W + "t"):
        first.remove(t)
    t = first.makeelement(W + "t", {"{http://www.w3.org/XML/1998/namespace}space": "preserve"})
    t.text = text
    first.append(t)


def strip_para_ids(el):
    """Bo w14:paraId / textId de ban nhan khong trung ma voi doan goc."""
    for node in el.iter():
        for key in list(node.attrib):
            if key.endswith("}paraId") or key.endswith("}textId"):
                del node.attrib[key]


def clone_para(ref, text):
    el = copy.deepcopy(ref._p)
    strip_para_ids(el)
    set_text(el, text)
    return el


def main():
    src = find_template()
    doc = Document(src)
    normalise_measures(doc)
    ps = doc.paragraphs
    body_ref, img_ref, cap_ref = ps[2], ps[3], ps[4]
    tt_ref, mt_ref, xl_ref = ps[18], ps[19], ps[20]
    assert body_ref.text.startswith("Khi truy cập") and cap_ref.text.startswith("Hình 3.4")
    assert tt_ref.text.startswith("Tình huống 4")

    ext = img_ref._p.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
    width = Emu(int(ext.get("cx")))

    # --- 1. mo ta + hinh: chen sau chu thich Hinh 3.4 -----------------------
    anchor = cap_ref._p
    for text, png, caption in MO_TA:
        path = os.path.join(BA_DIR, png)
        if not os.path.exists(path):
            raise SystemExit("thieu anh: " + path)
        for el in (clone_para(body_ref, text),):
            anchor.addnext(el)
            anchor = el
        img_p = doc.add_paragraph()
        apply_pPr(img_p, img_ref._p.find(W + "pPr"))
        img_p.add_run().add_picture(path, width=width)
        anchor.addnext(img_p._p)
        anchor = img_p._p
        cap = clone_para(cap_ref, caption)
        anchor.addnext(cap)
        anchor = cap

    # --- 2. bang 3.3 ---------------------------------------------------------
    table = doc.tables[0]
    last = table.rows[-1]._tr
    for row in QUY_TAC:
        new_tr = copy.deepcopy(last)
        strip_para_ids(new_tr)
        for tc, value in zip(new_tr.findall(W + "tc"), row):
            set_text(tc.find(W + "p"), value)
        last.addnext(new_tr)
        last = new_tr

    # --- 3. tinh huong ---------------------------------------------------------
    anchor = xl_ref._p
    for title, mota, xuly in TINH_HUONG:
        for ref, text in ((tt_ref, title), (mt_ref, mota), (xl_ref, xuly)):
            el = clone_para(ref, text)
            anchor.addnext(el)
            anchor = el

    doc.save(OUT)
    print("nguon:", src)
    print("docx:", OUT)


if __name__ == "__main__":
    main()
