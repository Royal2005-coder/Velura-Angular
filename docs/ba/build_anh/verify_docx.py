# -*- coding: utf-8 -*-
"""Kiem tra cau truc docx sau khi bo sung."""
import re
import sys
import zipfile
import os

from docx import Document

PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "Velura-3.1.3-Tim-kiem-bang-hinh-anh.docx")
d = Document(PATH)
for i, p in enumerate(d.paragraphs):
    flag = "IMG" if "graphicData" in p._p.xml else ""
    print(i, p.style.name, flag, p.text[:70].replace("\n", " "))
t = d.tables[0]
print("rows:", len(t.rows))
for r in t.rows:
    print([c.text[:18] for c in r.cells])
z = zipfile.ZipFile(PATH)
print([n for n in z.namelist() if "media" in n])
x = z.read("word/document.xml").decode("utf-8")
ids = re.findall(r'paraId="([0-9A-Fa-f]+)"', x)
print("paraId:", len(ids), "unique:", len(set(ids)))
sys.exit(0)
