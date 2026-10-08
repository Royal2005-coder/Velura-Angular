import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_uat_workbook():
    wb = openpyxl.Workbook()
    
    # ------------------ PALETTE & FONTS ------------------
    FONT_NAME = "Calibri"
    
    NAVY_DARK = "1E293B"    # Slate 800 - Primary Brand Navy
    GOLD_BROWN = "7D562D"   # Velura Gold Brown
    ACCENT_BLUE = "0284C7"  # Sky Blue for secondary
    BG_LIGHT = "F8FAFC"     # Off-white / light slate
    BORDER_LIGHT = "E2E8F0" # Slate 200 border
    
    STATUS_PASS_FILL = "DCFCE7"    # Green 100
    STATUS_PASS_FONT = "166534"    # Green 800
    STATUS_FAIL_FILL = "FEE2E2"    # Red 100
    STATUS_FAIL_FONT = "991B1B"    # Red 800
    STATUS_UNTESTED_FILL = "F1F5F9"# Slate 100
    STATUS_UNTESTED_FONT = "475569"# Slate 600

    thin_border = Border(
        left=Side(style='thin', color=BORDER_LIGHT),
        right=Side(style='thin', color=BORDER_LIGHT),
        top=Side(style='thin', color=BORDER_LIGHT),
        bottom=Side(style='thin', color=BORDER_LIGHT)
    )
    header_border = Border(
        left=Side(style='thin', color='475569'),
        right=Side(style='thin', color='475569'),
        top=Side(style='medium', color='0F172A'),
        bottom=Side(style='medium', color='0F172A')
    )

    # =========================================================================
    # TAB DATA DEFINITIONS
    # =========================================================================
    tabs_data = [
        {
            "sheet_name": "2.1 Mua Hàng E2E",
            "tester_code": "Tester 01",
            "tester_name": "Nguyễn Thị Ánh",
            "persona": "Khách Mua Sắm Trực Tuyến (Online Shopper)",
            "flow_title": "QUY TRÌNH 2.1: MUA HÀNG TRỰC TUYẾN E2E (TÌM KIẾM, BIẾN THỂ, GIỎ HÀNG, CHECKOUT STRIPE & COD)",
            "account_info": "Khách vãng lai (Guest) & Khách đăng nhập: shopper.uat@velura.test / Pass: Velura@2026",
            "url_access": "Storefront: http://localhost:4200 (hoặc https://royalai.dev)",
            "cases": [
                (
                    "TC-BUY-01",
                    "Tìm kiếm & Lọc sản phẩm",
                    "Khách truy cập trang chủ Velura Storefront",
                    "1. Nhập từ khóa 'áo sơ mi lụa' vào thanh tìm kiếm header.\n2. Áp dụng bộ lọc: Danh mục 'Áo', Khoảng giá '300.000đ - 1.000.000đ', Màu 'Trắng'.\n3. Kiểm tra danh sách kết quả hiển thị.",
                    "Từ khóa: 'áo sơ mi lụa' | Bộ lọc: Áo, 300k-1tr",
                    "Hệ thống trả về danh sách sản phẩm đúng tiêu chí lọc trong < 1s; hiển thị ảnh đại diện, tên, giá bán và badge giảm giá nếu có.",
                    "Untested", "Major"
                ),
                (
                    "TC-BUY-02",
                    "Tìm kiếm sản phẩm bằng hình ảnh AI (Visual Search)",
                    "User có sẵn ảnh chụp mẫu trang phục trong máy",
                    "1. Bấm vào icon máy ảnh trên thanh tìm kiếm header.\n2. Chọn tải lên file ảnh áo thời trang mẫu (PNG/JPG).\n3. Chờ AI xử lý phân tích hình ảnh và trả về danh sách gợi ý.",
                    "File: sample_blouse.jpg (< 5MB)",
                    "Hệ thống phân tích hình ảnh qua AI Worker, trả về các mẫu áo có kiểu dáng, màu sắc hoặc họa tiết tương đồng nhất trong catalog.",
                    "Untested", "Major"
                ),
                (
                    "TC-BUY-03",
                    "Chi tiết sản phẩm (PDP) - Phối hợp biến thể Màu/Size",
                    "User mở trang chi tiết một sản phẩm thời trang nhiều biến thể",
                    "1. Mở trang PDP `/products/ao-so-mi-lua-velura-01`.\n2. Lần lượt click chọn các ô Màu sắc: Trắng -> Đen -> Be.\n3. Chọn Kích thước: S -> M -> L.\n4. Quan sát hình ảnh gallery lớn và giá tiền tương ứng.",
                    "Sản phẩm: Áo Sơ Mi Lụa Cao Cấp (SP-001)",
                    "Ảnh sản phẩm chính lập tức đổi theo màu sắc vừa chọn; Giá tiền và thông tin tình trạng tồn kho cập nhật mượt mà, không bị giật lag.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BUY-04",
                    "Kiểm tra khóa thao tác khi biến thể HẾT HÀNG (Out of Stock)",
                    "Sản phẩm có ít nhất 1 biến thể có số lượng tồn kho = 0",
                    "1. Trên trang PDP, chọn biến thể Màu 'Đen' - Size 'XL' (tồn kho = 0).\n2. Kiểm tra giao diện hiển thị của nút Màu/Size và nút hành động.\n3. Thử nhấn vào nút 'Thêm vào giỏ' và 'Mua ngay'.",
                    "Biến thể: Đen - XL (Stock = 0)",
                    "Ô biến thể bị gạch chéo mờ (disabled); Cả hai nút 'Thêm vào giỏ' và 'Mua ngay' đều bị vô hiệu hóa (disabled), hiển thị badge 'Hết hàng'. Chặn hoàn toàn không cho mua.",
                    "Untested", "Blocker"
                ),
                (
                    "TC-BUY-05",
                    "Thêm nhanh vào giỏ hàng (Quick Add)",
                    "Sản phẩm còn hàng trong kho",
                    "1. Chọn biến thể còn hàng (Màu Trắng - Size M).\n2. Nhấn nút 'Thêm vào giỏ hàng'.\n3. Quan sát góc trên bên phải màn hình và icon giỏ hàng.",
                    "Biến thể: Trắng - M (Stock = 25)",
                    "Hiển thị Toast thông báo 'Đã thêm vào giỏ hàng thành công!'; Icon giỏ hàng tăng số lượng tức thì; Không bị chuyển hướng trang ngoài ý muốn.",
                    "Untested", "Major"
                ),
                (
                    "TC-BUY-06",
                    "Nút 'Mua ngay' chuyển thẳng Checkout (Buy Now)",
                    "Sản phẩm còn hàng trong kho",
                    "1. Chọn biến thể Màu Be - Size S.\n2. Nhấn nút 'Mua ngay'.\n3. Kiểm tra trang đích được chuyển đến và danh sách sản phẩm trong đơn.",
                    "Biến thể: Be - S | Số lượng: 1",
                    "Hệ thống chuyển thẳng lập tức sang trang `/checkout` (hoặc `/checkout/guest`); Đúng sản phẩm và biến thể vừa chọn được nạp sẵn vào tóm tắt đơn.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BUY-07",
                    "Quản lý Giỏ hàng `/cart` - Thay đổi biến thể & số lượng",
                    "Giỏ hàng đang có ít nhất 2 sản phẩm",
                    "1. Truy cập trang `/cart`.\n2. Đổi số lượng từ 1 lên 3 chiếc -> Kiểm tra tạm tính.\n3. Bấm vào đổi Size trực tiếp từ M sang L ngay trong giỏ hàng.\n4. Bấm xóa 1 sản phẩm khỏi giỏ.",
                    "Cart Items: Áo sơ mi (qty 1->3), Quần tây",
                    "Tổng tiền và phí vận chuyển tự động tính lại chuẩn xác theo công thức; Đổi size trực tiếp cập nhật ngay; Xóa sản phẩm hiển thị toast xác nhận và xóa khỏi danh sách.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BUY-08",
                    "Thanh toán Đơn hàng Khách vãng lai (Guest Checkout - COD)",
                    "Khách chưa đăng nhập tài khoản, có sản phẩm trong giỏ",
                    "1. Từ giỏ hàng bấm 'Tiến hành đặt hàng' -> Mở `/checkout/guest`.\n2. Điền họ tên, SĐT `0987654321`, địa chỉ TP.HCM.\n3. Chọn phương thức 'Thanh toán khi nhận hàng (COD)'.\n4. Nhấn 'Đặt hàng ngay'.",
                    "Họ tên: Lê Minh Anh | SĐT: 0987654321 | COD",
                    "Đặt hàng thành công; Chuyển hướng sang `/checkout/confirm`; Hiển thị Mã đơn hàng định dạng `VEL-2026-XXXX`; Hiển thị lời mời tạo tài khoản để tích điểm VIP.",
                    "Untested", "Blocker"
                ),
                (
                    "TC-BUY-09",
                    "Thanh toán Thẻ tín dụng Quốc tế Stripe (Member Checkout)",
                    "Khách đã đăng nhập tài khoản thành viên",
                    "1. Mở `/checkout/user` -> Chọn địa chỉ đã lưu trong sổ địa chỉ.\n2. Chọn hình thức 'Thanh toán qua thẻ (Stripe)'.\n3. Nhập thẻ test Stripe: Số `4242 4242 4242 4242`, CVC `123`, Hạn `12/28`.\n4. Bấm 'Thanh toán ngay'.",
                    "Thẻ Stripe Test: 4242424242424242 | Exp: 12/28",
                    "Cổng Stripe xác thực thành công trong 2 giây; Hệ thống ghi nhận trạng thái thanh toán `PAID`; Hiển thị màn hình đặt hàng thành công.",
                    "Untested", "Blocker"
                ),
                (
                    "TC-BUY-10",
                    "Email thông báo Đặt hàng & Hóa đơn tự động",
                    "Đơn hàng đặt thành công ở bước TC-BUY-08 hoặc TC-BUY-09",
                    "1. Kiểm tra hòm thư email của khách hàng đã nhập khi đặt hàng.\n2. Mở email xác nhận đơn hàng từ sender Velura Fashion.\n3. Kiểm tra chi tiết mã đơn, sản phẩm, địa chỉ và đường link tra cứu.",
                    "Email: shopper.uat@velura.test",
                    "Email gửi về trong vòng < 30 giây; Layout email chuyên nghiệp, hiển thị đúng mã đơn `VEL-2026-XXXX`, số tiền, link theo dõi đơn hoạt động tốt.",
                    "Untested", "Major"
                )
            ]
        },
        {
            "sheet_name": "2.2 Hủy Đơn & Đổi Trả E2E",
            "tester_code": "Tester 02",
            "tester_name": "Trần Quốc Bảo",
            "persona": "Khách Hậu Mãi & Chuyên Viên RMA CSKH",
            "flow_title": "QUY TRÌNH 2.2: HỦY ĐƠN HÀNG, YÊU CẦU ĐỔI TRẢ (RMA), UPLOAD ẢNH & CSKH DUYỆT THỦ TỤC",
            "account_info": "User: customer.rma@velura.test / Pass: Velura@2026 | Admin: cskh.rma@velura.vn / Pass: Admin@2026",
            "url_access": "Storefront: http://localhost:4200/account/returns | Admin: http://localhost:4201/returns",
            "cases": [
                (
                    "TC-RMA-01",
                    "Khách tự hủy đơn hàng Chờ xử lý (PENDING)",
                    "Tài khoản có đơn hàng vừa tạo ở trạng thái PENDING",
                    "1. Đăng nhập User -> Vào trang `/account/orders`.\n2. Chọn đơn hàng trạng thái 'Chờ xử lý' -> Bấm 'Xem chi tiết'.\n3. Nhấn nút 'Hủy đơn hàng' -> Chọn lý do 'Muốn đổi mẫu khác'.\n4. Xác nhận hủy đơn.",
                    "Đơn hàng PENDING vừa đặt | Lý do hủy",
                    "Đơn hàng chuyển trạng thái ngay sang `CANCELLED`; Hệ thống tự động hoàn lại số lượng tồn kho cho các sản phẩm trong đơn.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-02",
                    "Kiểm tra chặn hủy đơn khi đơn ĐÃ DUYỆT / ĐANG GIAO",
                    "Đơn hàng đã chuyển sang trạng thái CONFIRMED hoặc SHIPPING",
                    "1. Vào chi tiết đơn hàng đã xuất kho hoặc đang giao.\n2. Kiểm tra sự xuất hiện của nút 'Hủy đơn hàng'.",
                    "Đơn hàng trạng thái SHIPPING",
                    "Nút 'Hủy đơn' bị ẩn hoặc vô hiệu hóa; Hiển thị thông báo: 'Đơn hàng đang trên đường giao, vui lòng liên hệ hotline 1900-xxxx nếu cần hỗ trợ'.",
                    "Untested", "Major"
                ),
                (
                    "TC-RMA-03",
                    "Khởi tạo Yêu cầu Đổi trả (RMA) trên đơn đã giao",
                    "Đơn hàng đã ở trạng thái DELIVERED trong vòng 30 ngày",
                    "1. Vào `/account/orders` -> Chọn đơn đã giao thành công.\n2. Nhấn nút 'Yêu cầu đổi / trả' -> Mở trang form `/account/returns`.\n3. Chọn 1 sản phẩm trong đơn cần đổi trả.",
                    "Đơn hàng DELIVERED: Áo dạ tweed cao cấp (850.000đ)",
                    "Form đổi trả tải lên thông tin đơn hàng chuẩn xác; Hiển thị quy định 30 ngày đổi trả miễn phí.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-04",
                    "Tải lên hình ảnh bằng chứng sản phẩm lỗi (Upload ảnh RMA)",
                    "User có ảnh chụp chi tiết lỗi bung chỉ / rách vải của áo",
                    "1. Chọn lý do 'Sản phẩm lỗi kỹ thuật / bung chỉ'.\n2. Nhấn chọn tải lên 2 file ảnh chụp lỗi sản phẩm.\n3. Kiểm tra xem trước ảnh (Preview ảnh thumbnail) và thử bấm icon xóa 1 ảnh.\n4. Nhập ghi chú chi tiết mô tả lỗi.",
                    "2 file ảnh JPG lỗi (< 5MB/ảnh)",
                    "Ảnh tải lên hiển thị thumbnail xem trước sắc nét; Xóa ảnh hoạt động tốt; Chặn file > 5MB hoặc file sai định dạng.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-05",
                    "Chọn hình thức HOÀN TIỀN (Refund) & Cố định giá trị",
                    "Đang thao tác tại form tạo yêu cầu đổi trả",
                    "1. Chọn hình thức giải quyết: 'Hoàn tiền vào tài khoản'.\n2. Nhập thông tin STK Ngân hàng, Tên chủ thẻ, Tên ngân hàng.\n3. Kiểm tra ô 'Số tiền hoàn dự kiến'.",
                    "STK: 1903xxx Techcombank | Số tiền: 850.000đ",
                    "Hệ thống khóa cố định số tiền hoàn theo đúng giá trị thanh toán thực tế của sản phẩm trên hóa đơn; Khách hàng không thể sửa đổi số tiền.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-06",
                    "Chọn hình thức ĐỔI HÀNG (Exchange) - Đổi Size/Màu",
                    "Đang thao tác tại form tạo yêu cầu đổi trả",
                    "1. Chọn hình thức: 'Đổi sản phẩm / Đổi kích cỡ'.\n2. Chọn Biến thể mới muốn đổi: Màu 'Đen' - Size 'L'.\n3. Kiểm tra tồn kho của biến thể mới được chọn.\n4. Bấm 'Gửi yêu cầu đổi trả'.",
                    "Sản phẩm đổi: Áo dạ tweed Đen - Size L (còn hàng)",
                    "Yêu cầu đổi trả được tạo thành công với mã `RMA-2026-XXXX`; Trạng thái hiển thị `REQUESTED`; User nhận được email xác nhận.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-07",
                    "Admin CSKH tiếp nhận phiếu RMA & Chuyển 'CONTACTING'",
                    "Admin CSKH đăng nhập trang `/admin/returns`",
                    "1. Mở danh sách đổi trả -> Thấy phiếu RMA vừa tạo.\n2. Bấm nút Thao tác (icon bút chì) -> Chọn 'Liên hệ khách hàng'.\n3. Nhập ghi chú: 'Đã gọi khách xác nhận địa chỉ lấy hàng đổi'.",
                    "Phiếu RMA-2026-XXXX vừa tạo",
                    "Trạng thái phiếu đổi sang `CONTACTING`; Cột Thao tác đổi dropdown menu tương ứng; Không bị xô lệch layout nút bấm.",
                    "Untested", "Major"
                ),
                (
                    "TC-RMA-08",
                    "Admin CSKH phê duyệt Đổi hàng / Hoàn tiền (APPROVED)",
                    "Phiếu RMA đang ở trạng thái CONTACTING",
                    "1. Tại cột Thao tác, click mở Dropdown menu hành động.\n2. Chọn 'Phê duyệt đổi hàng' (hoặc 'Phê duyệt hoàn tiền').\n3. Nhập hướng dẫn gửi hàng cho khách: 'Shipper Velura sẽ đến thu hồi kiện hàng tận nơi trong 24h'.\n4. Xác nhận duyệt.",
                    "Ghi chú phê duyệt CSKH",
                    "Phiếu chuyển trạng thái `APPROVED`; Khách hàng kiểm tra giao diện Storefront `/account/returns` thấy ngay trạng thái đổi sang 'Đã được duyệt'.",
                    "Untested", "Critical"
                ),
                (
                    "TC-RMA-09",
                    "Admin CSKH từ chối yêu cầu RMA (REJECTED)",
                    "Có phiếu RMA vi phạm chính sách (quá hạn 30 ngày hoặc sản phẩm đã qua giặt tẩy)",
                    "1. Chọn phiếu RMA vi phạm -> Mở dropdown Thao tác.\n2. Chọn 'Từ chối đổi trả'.\n3. Thử để trống lý do bấm xác nhận -> Kiểm tra validation.\n4. Nhập lý do từ chối cụ thể (> 10 ký tự) và xác nhận.",
                    "Lý do: 'Sản phẩm đã bị cắt tem mác và qua sử dụng giặt tẩy'",
                    "Hệ thống bắt buộc nhập lý do từ chối >= 10 ký tự; Chuyển trạng thái sang `REJECTED`; Gửi email thông báo từ chối kèm lý do cho khách hàng.",
                    "Untested", "Major"
                ),
                (
                    "TC-RMA-10",
                    "Kho nhận hàng & Kiểm định QA (QA Inspection & Hoàn tiền)",
                    "Kiện hàng khách gửi về kho (Trạng thái RETURNING)",
                    "1. Thủ kho đăng nhập Admin -> Mở chi tiết phiếu RMA.\n2. Bấm 'Tiếp nhận kiện hàng' -> Chuyển sang `RECEIVED`.\n3. Tiến hành kiểm định chất lượng: Chọn 'QA Pass'.\n4. Bấm 'Hoàn tiền cho khách' (Stripe hoàn tự động hoặc xác nhận chuyển khoản COD).",
                    "Kiểm tra tem mác, độ nguyên vẹn vải",
                    "Trạng thái chuyển sang `REFUNDED` / `COMPLETED`; Đơn hàng kết thúc chu trình; Ghi log kiểm định chất lượng vào lịch sử RMA.",
                    "Untested", "Blocker"
                )
            ]
        },
        {
            "sheet_name": "2.3 AI Chatbot E2E",
            "tester_code": "Tester 03",
            "tester_name": "Lê Hoàng Châu",
            "persona": "Khách Hàng Trải Nghiệm AI Stylist Chatbot",
            "flow_title": "QUY TRÌNH 2.3: TRỢ LÝ THỜI TRANG AI CHATBOT (TƯ VẤN OUTFIT, GỢI Ý SP, GỬI ẢNH MẪU, GIỎ HÀNG)",
            "account_info": "Khách vãng lai & Thành viên: member.chatbot@velura.test / Pass: Velura@2026",
            "url_access": "Storefront: http://localhost:4200/chatbot (hoặc icon Chatbot góc phải màn hình)",
            "cases": [
                (
                    "TC-BOT-01",
                    "Khởi tạo phiên Chatbot & Lời chào AI Stylist",
                    "User truy cập trang `/chatbot` trên Storefront",
                    "1. Truy cập đường dẫn `/chatbot`.\n2. Quan sát khung chat chính giữa và thanh sidebar lịch sử trò chuyện bên trái.\n3. Kiểm tra tin nhắn chào mừng mặc định `localGreeting()`.",
                    "Trang `/chatbot`",
                    "Giao diện chatbot chuẩn phong cách thương hiệu Velura tải mượt mà; Hiển thị lời chào cá nhân hóa và các nút gợi ý câu hỏi nhanh (Quick Prompts).",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-02",
                    "Tư vấn phối đồ theo vóc dáng và dịp lễ (Styling Consultation)",
                    "Khung chat đang mở sẵn sàng nhập liệu",
                    "1. Nhập vào ô chat: 'Tôi cao 1m58 nặng 48kg, hãy tư vấn cho tôi 1 set đồ đi tiệc cưới ngoài trời mùa thu sang trọng'.\n2. Nhấn Enter hoặc icon gửi tin nhắn.\n3. Quan sát chỉ báo AI đang phản hồi (Typing indicator).",
                    "Nội dung prompt tư vấn vóc dáng 1m58",
                    "Hiển thị animation đang soạn câu trả lời; Sau < 3s, AI Stylist phản hồi tư vấn chi tiết về form dáng, chất liệu lụa/tweed và bảng màu phù hợp tone da.",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-03",
                    "Hiển thị Thẻ Sản Phẩm Tương Tác (Product Cards) trong hội thoại",
                    "AI phản hồi câu hỏi tư vấn phối đồ",
                    "1. Quan sát phần bên dưới câu trả lời của AI Stylist.\n2. Kiểm tra các thẻ sản phẩm đính kèm (`productsFor`).\n3. Kiểm tra ảnh sản phẩm, tên, giá tiền VND format chuẩn, và tình trạng tồn kho.",
                    "Dữ liệu catalog sản phẩm trả về từ API",
                    "Các sản phẩm gợi ý hiển thị dạng carousel/grid thẻ bài đẹp mắt; Mỗi thẻ có đầy đủ ảnh, tên, giá format tiền tệ chuẩn (VD: 650.000 ₫) và badge còn hàng.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BOT-04",
                    "Thêm nhanh vào giỏ hàng (Add to Cart) trực tiếp từ Card Chatbot",
                    "Thẻ sản phẩm hiển thị trong khung chat",
                    "1. Bấm nút 'Thêm vào giỏ' ngay trên thẻ sản phẩm trong khung chat.\n2. Quan sát icon giỏ hàng trên header website.\n3. Bấm mở xem giỏ hàng mini hoặc trang `/cart`.",
                    "Sản phẩm trên card: Đầm xòe lụa satin",
                    "Sản phẩm được nạp ngay lập tức vào giỏ hàng (`CartStore`); Header giỏ hàng tăng số lượng; Hiển thị thông báo thành công trong khung chat.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BOT-05",
                    "Tải ảnh trang phục lên để Bot phân tích (Visual Style Consultation)",
                    "User có ảnh mẫu trang phục hoặc phối đồ trong máy",
                    "1. Bấm vào icon đính kèm hình ảnh (nút kẹp giấy / máy ảnh).\n2. Chọn 1 file ảnh mẫu outfit từ máy tính.\n3. Kiểm tra preview ảnh xem trước phía trên ô nhập liệu.\n4. Nhập kèm: 'Bộ này nên phối với phụ kiện và túi xách màu gì?' -> Gửi.",
                    "File: wedding_dress_sample.png (< 5MB)",
                    "Ảnh hiển thị preview xem trước có nút X để gỡ nếu muốn; Khi bấm gửi, ảnh đính kèm được tải lên và Bot phân tích màu sắc, đưa ra lời khuyên phụ kiện cực kỳ chuẩn xác.",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-06",
                    "Hỏi đáp Chính sách thương hiệu & Đổi trả 30 ngày (FAQ Knowledge)",
                    "Khung chat đang hoạt động",
                    "1. Nhập câu hỏi: 'Chính sách đổi trả của Velura như thế nào? Có mất phí ship không?'.\n2. Gửi tin nhắn và kiểm tra nội dung AI trả về.",
                    "Câu hỏi chính sách đổi trả / vận chuyển",
                    "Bot trả lời chính xác theo cẩm nang thương hiệu: Hỗ trợ đổi trả miễn phí trong 30 ngày tận nơi, hoàn tiền 100% nếu có lỗi từ nhà sản xuất.",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-07",
                    "Nhận diện ý định tra cứu tình trạng đơn hàng qua Chatbot",
                    "User có mã đơn hàng cụ thể",
                    "1. Nhập vào ô chat: 'Kiểm tra giúp tôi tình trạng đơn hàng VEL-2026-1024'.\n2. Gửi tin nhắn và kiểm tra phản hồi.",
                    "Mã đơn hàng: VEL-2026-1024",
                    "Bot nhận diện đúng Intent tra cứu đơn hàng; Trả về trạng thái hiện tại của đơn kèm đường link trực tiếp dẫn tới trang `/account/track` hoặc `/account/orders/VEL-2026-1024`.",
                    "Untested", "Critical"
                ),
                (
                    "TC-BOT-08",
                    "Gợi ý bài viết Tạp chí phong cách (Blog Articles Card)",
                    "Hỏi về xu hướng thời trang mới nhất",
                    "1. Nhập câu hỏi: 'Xu hướng thời trang công sở mùa thu đông năm nay có gì nổi bật?'.\n2. Gửi tin nhắn và quan sát kết quả.",
                    "Chủ đề xu hướng thời trang mùa thu",
                    "Bên cạnh câu trả lời văn bản, Bot đính kèm thẻ bài viết Blog (`blogsFor`) với ảnh bìa, tiêu đề và link click mở bài viết `/blog/xu-huong-thoi-trang-cong-so-thu-dong-2026`.",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-09",
                    "Quản lý Phiên trò chuyện (New Chat, Chọn phiên cũ, Xóa phiên)",
                    "User đã có ít nhất 2 phiên chat trước đó",
                    "1. Bấm nút 'Cuộc trò chuyện mới' (New Chat) -> Kiểm tra màn hình làm mới.\n2. Bấm vào một phiên chat cũ trong Sidebar danh sách -> Kiểm tra tải lại lịch sử tin nhắn.\n3. Bấm icon thùng rác tại một phiên chat để xóa -> Kiểm tra phiên bị xóa.",
                    "Sidebar danh sách Chat Sessions",
                    "Tạo phiên mới tức thì; Tải lại đúng toàn bộ lịch sử tin nhắn của phiên cũ đã chọn; Xóa phiên thành công và cập nhật lại danh sách sidebar.",
                    "Untested", "Major"
                ),
                (
                    "TC-BOT-10",
                    "Xử lý ngoại lệ, Ngắt kết nối & Phản hồi an toàn (Fallback & Robustness)",
                    "Thử nghiệm gửi câu hỏi không liên quan hoặc ngắt kết nối mạng",
                    "1. Nhập câu hỏi không thuộc phạm vi thời trang (VD: 'Giải phương trình toán học x^2 + 5x + 6 = 0').\n2. Thử tắt mạng tạm thời và bấm gửi tin nhắn.\n3. Bật mạng lại và bấm nút 'Thử lại' nếu có lỗi.",
                    "Prompt ngoài phạm vi & kịch bản mất kết nối",
                    "Với câu hỏi ngoài phạm vi: Bot lịch sự từ chối và hướng sự chú ý về thời trang Velura; Khi mất mạng: Hiển thị thông báo kết nối lỗi thân thiện, không crash ứng dụng.",
                    "Untested", "Critical"
                )
            ]
        },
        {
            "sheet_name": "2.4 Quản Lý & Tra Cứu Đơn E2E",
            "tester_code": "Tester 04",
            "tester_name": "Phạm Văn Dũng",
            "persona": "Khách Tra Cứu Đơn Hàng & Điều Phối Kho Fulfillment",
            "flow_title": "QUY TRÌNH 2.4: QUẢN LÝ, TRA CỨU ĐƠN HÀNG (MEMBER & GUEST) VÀ ĐIỀU PHỐI VẬN CHUYỂN 3PL",
            "account_info": "Guest: SĐT 0912345678 / Đơn VEL-2026-8899 | Member: member.tracking@velura.test | Admin Kho: warehouse.ops@velura.vn",
            "url_access": "Tra cứu Guest: http://localhost:4200/account/track | Admin Đơn: http://localhost:4201/orders",
            "cases": [
                (
                    "TC-ORD-01",
                    "Tra cứu đơn hàng Khách vãng lai (Guest Tracking)",
                    "Khách chưa đăng nhập tài khoản, có đơn hàng đã đặt thành công",
                    "1. Truy cập trang Tra cứu đơn hàng `/account/track` (hoặc `/guest/orders`).\n2. Nhập Mã đơn hàng: `VEL-2026-8899`.\n3. Nhập Số điện thoại đặt hàng: `0912345678`.\n4. Bấm nút 'Tra cứu đơn hàng'.",
                    "Mã đơn: VEL-2026-8899 | SĐT: 0912345678",
                    "Hệ thống xác thực thông tin hợp lệ; Hiển thị màn hình chi tiết đơn hàng khách vãng lai mà không yêu cầu đăng nhập mật khẩu.",
                    "Untested", "Critical"
                ),
                (
                    "TC-ORD-02",
                    "Hiển thị Chi tiết Đơn hàng & Tiến trình Giao nhận (Timeline Tracking)",
                    "Màn hình chi tiết đơn hàng khách vãng lai đã hiển thị",
                    "1. Kiểm tra danh sách sản phẩm (ảnh, tên biến thể, số lượng, đơn giá).\n2. Kiểm tra thanh timeline tiến độ giao nhận 5 bước: Đặt hàng -> Xác nhận -> Đóng gói -> Đang giao -> Đã giao.\n3. Kiểm tra thông tin đơn vị vận chuyển (ViettelPost / GHN) và mã vận đơn tracking.",
                    "Đơn hàng VEL-2026-8899",
                    "Thông tin đơn hàng hiển thị trực quan, đúng sản phẩm; Thanh tiến trình thể hiện đúng vị trí trạng thái hiện tại của đơn hàng theo thời gian thực.",
                    "Untested", "Critical"
                ),
                (
                    "TC-ORD-03",
                    "Validation Tra cứu đơn sai thông tin (Bảo mật thông tin)",
                    "Truy cập trang `/account/track`",
                    "1. Nhập Mã đơn đúng `VEL-2026-8899` nhưng nhập sai SĐT `0999999999`.\n2. Bấm 'Tra cứu đơn hàng'.\n3. Nhập mã đơn không tồn tại `VEL-9999-0000` kèm SĐT đúng.",
                    "Mã đơn hoặc SĐT không khớp",
                    "Hiển thị thông báo lỗi rõ ràng: 'Không tìm thấy đơn hàng phù hợp với thông tin đã cung cấp. Vui lòng kiểm tra lại mã đơn và số điện thoại'; Không tiết lộ dữ liệu của người khác.",
                    "Untested", "Major"
                ),
                (
                    "TC-ORD-04",
                    "Chuyển đổi Khách vãng lai sang Thành viên (Claim Account)",
                    "Khách vãng lai đang xem đơn hàng tại `/guest/orders`",
                    "1. Quan sát banner gợi ý: 'Tạo tài khoản thành viên để tích điểm thưởng và quản lý đơn dễ dàng'.\n2. Nhấn nút 'Kích hoạt tài khoản ngay' -> Mở `/account/claim`.\n3. Thiết lập mật khẩu mới và xác nhận.",
                    "Email/SĐT của đơn hàng vãng lai",
                    "Tài khoản thành viên được khởi tạo thành công; Đơn hàng vãng lai được tự động đồng bộ gán vào tài khoản mới tạo; Khách được nâng hạng VIP tích điểm.",
                    "Untested", "Major"
                ),
                (
                    "TC-ORD-05",
                    "Quản lý Đơn hàng Thành viên (`/account/orders`) theo Tab phân loại",
                    "Khách hàng đăng nhập tài khoản thành viên",
                    "1. Đăng nhập tài khoản -> Vào mục 'Đơn hàng của tôi' (`/account/orders`).\n2. Lần lượt click chuyển qua các tab trạng thái: 'Tất cả', 'Chờ xác nhận', 'Đang giao', 'Đã giao', 'Đã hủy'.\n3. Kiểm tra số lượng và danh sách đơn tương ứng mỗi tab.",
                    "Tài khoản member.tracking@velura.test",
                    "Danh sách đơn hàng lọc mượt mà theo từng tab trạng thái; Mỗi đơn hàng hiển thị tóm tắt mã đơn, ngày đặt, tổng tiền, nút 'Xem chi tiết' và nút 'Mua lại'.",
                    "Untested", "Critical"
                ),
                (
                    "TC-ORD-06",
                    "Chi tiết Đơn hàng Thành viên & Tải Hóa đơn Điện tử VAT",
                    "Đang ở trang danh sách đơn hàng thành viên",
                    "1. Bấm 'Xem chi tiết' một đơn hàng đã giao thành công.\n2. Kiểm tra thông tin địa chỉ giao hàng, phương thức thanh toán, chi tiết giảm giá voucher.\n3. Nhấn nút 'Tải hóa đơn VAT' (hoặc In phiếu xuất).",
                    "Đơn hàng thành công có yêu cầu VAT",
                    "Hiển thị đầy đủ thông tin xuất hóa đơn VAT công ty; Tải xuống file hóa đơn PDF hoặc mở cửa sổ in ấn chuẩn layout chuyên nghiệp.",
                    "Untested", "Major"
                ),
                (
                    "TC-ORD-07",
                    "Admin Vận hành Lọc & Tìm kiếm Đơn hàng toàn hệ thống (`/admin/orders`)",
                    "Admin Kho / Điều phối viên đăng nhập trang quản trị",
                    "1. Mở trang `/admin/orders`.\n2. Dùng bộ lọc: Lọc đơn theo khoảng ngày tạo, theo trạng thái `pending`, theo phương thức `COD`.\n3. Tìm kiếm theo mã đơn hoặc số điện thoại khách hàng.",
                    "Bộ lọc Admin Orders",
                    "Bảng dữ liệu phản hồi nhanh (< 500ms); Hiển thị đúng số lượng đơn theo điều kiện lọc; Có phân trang rõ ràng nếu danh sách dài.",
                    "Untested", "Critical"
                ),
                (
                    "TC-ORD-08",
                    "Nhận diện & Ghi Log xác nhận đơn hàng COD có rủi ro",
                    "Đơn hàng COD giá trị cao (> 2.000.000đ) hoặc thông tin địa chỉ chưa rõ ràng",
                    "1. Tại danh sách đơn hàng Admin, thấy đơn có gắn cờ cảnh báo rủi ro.\n2. Nhân viên gọi điện cho khách -> Bấm nút 'Ghi nhận cuộc gọi'.\n3. Chọn kết quả: 'Khách xác nhận lấy hàng' -> Bấm duyệt đơn.",
                    "Đơn COD giá trị 2.500.000đ",
                    "Hệ thống lưu lại log cuộc gọi kèm tên nhân viên và thời gian; Trạng thái đơn hàng chuyển từ `PENDING` sang `CONFIRMED`.",
                    "Untested", "Major"
                ),
                (
                    "TC-ORD-09",
                    "Đóng gói, Xuất kho & Gán Mã Vận Đơn 3PL (Fulfillment & Shipping)",
                    "Đơn hàng đã ở trạng thái CONFIRMED",
                    "1. Mở chi tiết đơn hàng trong Admin -> Bấm 'Soạn hàng & Xuất kho'.\n2. Chọn Đơn vị vận chuyển: 'Viettel Post' (hoặc GHN/GHTK).\n3. Nhập Mã vận đơn: `VTP-883921094`.\n4. Bấm 'Bàn giao vận chuyển'.",
                    "Đơn vị: Viettel Post | Mã vận đơn: VTP-883921094",
                    "Trạng thái đơn hàng chuyển sang `SHIPPING`; Mã vận đơn 3PL được đồng bộ ngay lập tức sang màn hình tra cứu của khách hàng (cả Guest và Member).",
                    "Untested", "Blocker"
                ),
                (
                    "TC-ORD-10",
                    "Cập nhật Trạng thái Giao hàng Thành công (DELIVERED) & Mở tính năng Review",
                    "Đơn hàng đang ở trạng thái SHIPPING",
                    "1. Khi nhận webhook 3PL báo giao thành công (hoặc Admin cập nhật thủ công sang 'Đã giao hàng').\n2. Kiểm tra trạng thái đơn đổi thành `DELIVERED`.\n3. Khách hàng vào trang đơn hàng trên Storefront kiểm tra.",
                    "Cập nhật trạng thái DELIVERED",
                    "Đơn chuyển sang `DELIVERED`; Giao diện khách hàng xuất hiện nút 'Đánh giá sản phẩm' và nút 'Yêu cầu đổi / trả' trong thời hạn 30 ngày.",
                    "Untested", "Blocker"
                )
            ]
        },
        {
            "sheet_name": "2.5 Dashboard & Phân Tích E2E",
            "tester_code": "Tester 05",
            "tester_name": "Đặng Thị Mai",
            "persona": "Giám Đốc Kinh Doanh & Phân Tích Dữ Liệu (Business Analyst)",
            "flow_title": "QUY TRÌNH 2.5: DASHBOARD QUẢN TRỊ, PHÂN TÍCH DỮ LIỆU KINH DOANH, CHỈ SỐ VẬN HÀNH & SLA",
            "account_info": "Admin Quản trị / Phân tích: admin.analyst@velura.vn / Pass: Admin@2026 (Role: super_admin / admin_viewer)",
            "url_access": "Admin Backoffice: http://localhost:4201/dashboard",
            "cases": [
                (
                    "TC-DASH-01",
                    "Truy cập Dashboard & Kiểm tra Phân quyền RBAC",
                    "Đăng nhập tài khoản có quyền truy cập báo cáo quản trị",
                    "1. Đăng nhập trang Admin `/admin/login`.\n2. Kiểm tra chuyển hướng mặc định vào `/admin/dashboard`.\n3. Kiểm tra thanh tiêu đề, menu điều hướng và tên tài khoản đăng nhập.",
                    "Tài khoản admin.analyst@velura.vn",
                    "Trang Dashboard hiển thị đầy đủ, không bị lỗi 403 Forbidden; Giao diện chuẩn phong cách Velura Admin Dark/Navy chuyên nghiệp.",
                    "Untested", "Critical"
                ),
                (
                    "TC-DASH-02",
                    "Kiểm tra Tab VẬN HÀNH (Operations Tab) - 6 Thẻ Cảnh Báo Khẩn",
                    "Đang ở trang Dashboard, chọn tab 'Vận hành'",
                    "1. Chọn tab 'Vận hành' (Operations).\n2. Kiểm tra hiển thị của 6 thẻ cảnh báo cốt lõi:\n   + Đơn chờ duyệt (pendingOrders)\n   + Lỗi thanh toán (paymentErrors)\n   + Yêu cầu đổi trả đang mở (openReturns)\n   + Ticket hỗ trợ mở (openSupportTickets)\n   + Sản phẩm tồn kho thấp (lowStockProducts)\n   + Đánh giá khẩn cấp cần duyệt (urgentReviews).",
                    "Tab Vận hành Dashboard",
                    "Tất cả 6 thẻ cảnh báo đều hiển thị số liệu thực tế; Các thẻ có số cảnh báo > 0 được làm nổi bật với màu sắc cảnh báo trực quan (Đỏ/Cam).",
                    "Untested", "Critical"
                ),
                (
                    "TC-DASH-03",
                    "Thao tác Điều hướng nhanh từ Thẻ Cảnh Báo Vận Hành (Click-to-Action)",
                    "Thẻ cảnh báo có số lượng phát sinh",
                    "1. Bấm trực tiếp vào thẻ 'Yêu cầu đổi trả đang mở' -> Kiểm tra trang đích.\n2. Quay lại Dashboard, bấm vào thẻ 'Đơn hàng chờ duyệt' -> Kiểm tra trang đích.",
                    "Click vào các KPI Card",
                    "Click thẻ Đổi trả mở ngay trang `/admin/returns` kèm bộ lọc trạng thái chờ duyệt; Click thẻ Đơn hàng mở ngay `/admin/orders` lọc đơn `pending`.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-04",
                    "Kiểm tra Tab KINH DOANH (Business Tab) - Các Chỉ Số Tài Chính Cốt Lõi",
                    "Chuyển sang tab 'Kinh doanh' (Business)",
                    "1. Bấm chọn tab 'Kinh doanh'.\n2. Kiểm tra các chỉ số tài chính trọng yếu:\n   + Doanh thu thuần (Revenue VNĐ)\n   + Tổng đơn hàng thành công (Order count)\n   + Giá trị đơn hàng trung bình AOV (Average Order Value)\n   + Tỷ lệ hoàn thành đơn (Completion rate %)\n   + Doanh thu từ khuyến mãi (Promotion revenue & share).",
                    "Tab Kinh doanh Dashboard",
                    "Các con số tài chính được định dạng tiền tệ VNĐ chuẩn xác, tỷ lệ % hiển thị rõ ràng; Số liệu khớp với tổng các đơn hàng thành công trên hệ thống.",
                    "Untested", "Critical"
                ),
                (
                    "TC-DASH-05",
                    "Bộ lọc Khoảng Thời Gian (Time Range Filter: Ngày, Tuần, Tháng, Năm)",
                    "Đang xem số liệu tại Tab Kinh doanh",
                    "1. Lần lượt click chọn bộ lọc thời gian: 'Hôm nay' (day), '7 ngày gần nhất' (week), '30 ngày gần nhất' (month), 'Năm nay' (year).\n2. Quan sát nhãn thời gian `periodLabel` và các con số doanh thu.",
                    "Bộ lọc: day, week, month, year",
                    "Số liệu doanh thu và số đơn cập nhật tức thì theo đúng khoảng thời gian được chọn; Không xảy ra hiện tượng đứng trang hay tải lại toàn bộ giao diện.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-06",
                    "Biểu đồ Xu hướng Doanh thu (Revenue Trend Chart)",
                    "Bộ lọc thời gian đang chọn '30 ngày gần nhất'",
                    "1. Quan sát biểu đồ đường xu hướng doanh thu.\n2. Rê chuột (hover) vào các điểm mốc ngày trên đường biểu đồ.\n3. Kiểm tra tooltip hiển thị số tiền và ngày tương ứng.",
                    "Biểu đồ doanh thu 30 ngày",
                    "Biểu đồ vẽ mượt mà, trục tung thể hiện số tiền VNĐ, trục hoành thể hiện ngày tháng; Tooltip khi rê chuột hiển thị thông số chi tiết chuẩn xác.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-07",
                    "Phân tích Giọng nói Khách hàng & Chất lượng dịch vụ (Voice of Customer & CSAT)",
                    "Khu vực báo cáo phản hồi khách hàng trong Dashboard",
                    "1. Kiểm tra tỷ lệ đơn hàng được khách đánh giá (`coveragePct`).\n2. Kiểm tra điểm hài lòng dịch vụ CSKH trung bình (`csatAvg`).\n3. Kiểm tra tỷ lệ đổi trả (`returnRatePct`) và danh sách top lý do đổi trả hàng đầu (`returnReasons`).",
                    "Khối phân tích CSAT & RMA",
                    "Hiển thị bảng tỷ lệ đổi trả và xếp hạng các nguyên nhân phổ biến (chật size, lỗi vải...); Giúp ban giám đốc nhận diện nhanh vấn đề chất lượng sản phẩm.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-08",
                    "Phân tích Điểm nghẽn Đơn hàng & Tỷ lệ Hủy (Order Friction & Cancellation)",
                    "Khu vực báo cáo vận hành đơn hàng",
                    "1. Kiểm tra tổng số đơn hàng bị hủy (`cancelledOrders`).\n2. Kiểm tra số đơn giao không thành công (`failedDelivery`).\n3. Xem biểu đồ tròn phân tích các nguyên nhân hủy đơn chính (`cancelReasons`).",
                    "Báo cáo Order Friction",
                    "Hiển thị rõ ràng các lý do khách hủy đơn (đổi ý, đặt trùng, thời gian giao lâu); Phản ánh đúng dữ liệu từ các đơn hủy thực tế.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-09",
                    "Kiểm tra 9 Khối Quản trị Phân tích Chuyên sâu (AD_DB_01 -> AD_DB_09)",
                    "Cuộn xuống khu vực các khối báo cáo chuyên đề",
                    "1. Kiểm tra sự hiện diện của 9 khối phân tích: Hiệu quả bán hàng, Điểm nghẽn SLA, Sản phẩm bán chạy & than phiền, Đơn chưa đánh giá, CSAT CSKH, Voucher, Đổi trả theo dòng SP, v.v.\n2. Thử thao tác mở rộng / thu gọn các khối thông tin.",
                    "9 Khối quản trị AD_DB_01 đến AD_DB_09",
                    "Layout các khối phân tích trình bày ngay ngắn, mạch lạc theo chuẩn Dashboard doanh nghiệp; Cho phép quản trị viên xem chi tiết từng mảng hoạt động.",
                    "Untested", "Major"
                ),
                (
                    "TC-DASH-10",
                    "Tự động Đồng bộ Số liệu Mới (Realtime Data Refresh)",
                    "Mở song song 2 màn hình (Storefront đặt đơn và Admin Dashboard)",
                    "1. Khi Tester 01 vừa đặt xong 1 đơn hàng mới trên Storefront.\n2. Tại màn hình Dashboard Admin, bấm nút 'Làm mới dữ liệu' (Refresh icon).\n3. Kiểm tra số lượng đơn chờ duyệt và doanh thu tăng lên tương ứng.",
                    "Đơn hàng mới phát sinh từ Tester 01",
                    "Số liệu trên Dashboard cập nhật ngay lập tức sau khi nhấn làm mới; Thẻ 'Đơn chờ duyệt' tăng thêm 1 đơn, doanh thu ghi nhận chuẩn xác không cần reload trang.",
                    "Untested", "Critical"
                )
            ]
        },
        {
            "sheet_name": "2.6 Khuyến Mãi & Đồng Bộ E2E",
            "tester_code": "Tester 06",
            "tester_name": "Vũ Minh Tuấn",
            "persona": "Trưởng Phòng Marketing & Khách Hàng Săn Ưu Đãi",
            "flow_title": "QUY TRÌNH 2.6: QUẢN TRỊ KHUYẾN MÃI ADMIN, ĐỒNG BỘ STOREFRONT VÀ ÁP DỤNG VOUCHER TẠI CHECKOUT",
            "account_info": "Admin MKT: admin.marketing@velura.vn / Pass: Admin@2026 | User: member.promo@velura.test / Pass: Velura@2026",
            "url_access": "Admin: http://localhost:4201/promotions | User: http://localhost:4200/offers & /checkout",
            "cases": [
                (
                    "TC-PROMO-01",
                    "Admin Tạo Chiến Dịch Khuyến Mãi Mới (Campaign Creation)",
                    "Admin Marketing đăng nhập vào trang `/admin/promotions`",
                    "1. Mở tab 'Chiến dịch' (Campaigns) -> Bấm 'Thêm chiến dịch mới'.\n2. Nhập Tên chiến dịch: 'Lễ Hội Thời Trang Mùa Thu Velura 2026'.\n3. Chọn Thời gian: Từ ngày hiện tại đến hết tháng.\n4. Nhập mô tả chiến dịch và bấm 'Lưu chiến dịch'.",
                    "Chiến dịch: Lễ Hội Mùa Thu 2026",
                    "Chiến dịch mới được tạo thành công, xuất hiện ở đầu danh sách chiến dịch đang hoạt động; Ghi nhận nhật ký tạo chiến dịch vào audit log.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-02",
                    "Admin Tạo Mã Giảm Giá Voucher Theo % (Percentage Voucher)",
                    "Đang ở trang quản lý khuyến mãi `/admin/promotions`",
                    "1. Mở tab 'Mã giảm giá' (Vouchers) -> Bấm 'Tạo mã voucher'.\n2. Nhập Mã: `VELURA20` | Loại: Giảm 20% giá trị đơn hàng.\n3. Thiết lập Ràng buộc: Đơn tối thiểu `500.000đ`, Giảm tối đa `100.000đ`.\n4. Giới hạn: 500 lượt dùng, mỗi user dùng tối đa 1 lần.\n5. Bấm 'Lưu voucher'.",
                    "Mã: VELURA20 | Giảm: 20% | Min: 500k | Max giảm: 100k",
                    "Voucher tạo thành công, hiển thị trạng thái 'Đang hoạt động'; Ràng buộc giá trị đơn tối thiểu và mức giảm tối đa được lưu đúng cơ sở dữ liệu.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-03",
                    "Admin Cấu hình Voucher Độc Quyền Sinh Nhật Khách Hàng (Birthday Voucher)",
                    "Tạo mã ưu đãi dành riêng cho chương trình quà sinh nhật",
                    "1. Bấm 'Tạo mã voucher' -> Nhập Mã: `HPBD2026`.\n2. Loại giảm giá: Giảm tiền mặt cố định `150.000đ`.\n3. Đơn hàng tối thiểu: `600.000đ`.\n4. Đánh dấu tích chọn: 'Áp dụng cho chương trình Quà tặng Sinh nhật thành viên'.\n5. Lưu voucher.",
                    "Mã: HPBD2026 | Giảm: 150.000đ | Min order: 600.000đ",
                    "Hệ thống lưu cấu hình voucher sinh nhật; Tự động kích hoạt cơ chế đồng bộ cho các khách hàng có sinh nhật trong tháng.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-04",
                    "Admin Bật / Tắt Trạng Thái & Kiểm tra Nhật Ký Khuyến Mãi (Audit Log)",
                    "Danh sách voucher đã tạo",
                    "1. Tìm mã `VELURA20` -> Bấm nút gạt Tắt trạng thái kích hoạt.\n2. Kiểm tra voucher chuyển sang trạng thái 'Tạm ngưng'.\n3. Chuyển sang tab 'Nhật ký' (Audit logs) của phân hệ khuyến mãi.\n4. Kiểm tra lịch sử ghi lại hành động của admin.",
                    "Voucher VELURA20",
                    "Trạng thái voucher cập nhật tức thì; Tab Nhật ký ghi lại chính xác: 'Admin admin.marketing@velura.vn đã vô hiệu hóa voucher VELURA20 lúc HH:mm:ss'.",
                    "Untested", "Major"
                ),
                (
                    "TC-PROMO-05",
                    "Đồng bộ Chiến dịch & Voucher sang Trang Ưu Đãi Storefront (`/offers`)",
                    "Khách hàng truy cập Storefront",
                    "1. Bật lại voucher `VELURA20` bên Admin.\n2. Mở trình duyệt khách hàng truy cập trang `/offers`.\n3. Quan sát các banner chiến dịch và thẻ mã giảm giá hiển thị.",
                    "Trang Storefront `/offers`",
                    "Các voucher công khai đang hoạt động bên Admin lập tức đồng bộ hiển thị trên trang `/offers`; Khách có thể bấm nút 'Sao chép mã' để lưu vào bộ nhớ tạm.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-06",
                    "Đồng bộ Tự Động Voucher Sinh Nhật vào Ví Thành Viên (`/account/vouchers`)",
                    "Khách hàng có thông tin ngày sinh nhật trong tháng hiện tại",
                    "1. Đăng nhập tài khoản `member.promo@velura.test` (sinh nhật trong tháng).\n2. Truy cập trang cá nhân -> Mở 'Ví Voucher' (`/account/vouchers`).\n3. Kiểm tra danh sách voucher cá nhân.",
                    "Tài khoản có sinh nhật trong tháng",
                    "Mã ưu đãi sinh nhật `HPBD2026` tự động xuất hiện trong ví với badge 'Quà tặng Sinh nhật'; Hiển thị số tiền giảm 150.000đ và thời hạn sử dụng đến hết tháng.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-07",
                    "Áp dụng Voucher Thành Công tại màn hình Thanh toán Checkout (`/checkout`)",
                    "Khách hàng có giỏ hàng thỏa mãn điều kiện áp dụng",
                    "1. Mua sản phẩm có tổng giá trị `800.000đ` (> 600.000đ).\n2. Tiến hành vào trang `/checkout`.\n3. Tại ô 'Mã giảm giá', nhập mã `HPBD2026` -> Nhấn nút 'Áp dụng'.\n4. Quan sát bảng tóm tắt chi phí đơn hàng.",
                    "Đơn hàng 800.000đ | Mã: HPBD2026",
                    "Hệ thống xác thực mã hợp lệ; Dòng 'Giảm giá voucher' hiển thị trừ `-150.000đ`; Tổng tiền thanh toán giảm chính xác còn `650.000đ`.",
                    "Untested", "Blocker"
                ),
                (
                    "TC-PROMO-08",
                    "Kiểm tra Chặn Áp Dụng khi Đơn Chưa Đạt Giá Trị Tối Thiểu (Min Order Value)",
                    "Giỏ hàng có giá trị nhỏ hơn điều kiện voucher",
                    "1. Giỏ hàng hiện chỉ có 1 sản phẩm trị giá `400.000đ` (< 600.000đ).\n2. Vào trang `/checkout`, nhập mã `HPBD2026` -> Bấm 'Áp dụng'.\n3. Quan sát thông báo của hệ thống.",
                    "Đơn hàng 400.000đ | Mã: HPBD2026 (yêu cầu min 600k)",
                    "Hệ thống từ chối áp dụng và hiển thị thông báo lỗi rõ ràng: 'Đơn hàng tối thiểu phải từ 600.000đ để áp dụng voucher này'; Không giảm giá.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-09",
                    "Kiểm tra Chặn Áp Dụng Mã Hết Hạn hoặc Đã Bị Vô Hiệu Hóa",
                    "Mã voucher đã hết hạn hoặc bị Admin tắt",
                    "1. Tại trang Checkout, nhập mã voucher đã hết hạn hoặc mã không tồn tại `EXPIRED2025`.\n2. Thử nhập mã `VELURA20` khi đã bị admin tắt kích hoạt.\n3. Bấm 'Áp dụng'.",
                    "Mã voucher không hợp lệ / hết hạn",
                    "Hệ thống thông báo: 'Mã giảm giá không tồn tại hoặc đã hết hiệu lực'; Không cho phép trừ tiền gian lận.",
                    "Untested", "Critical"
                ),
                (
                    "TC-PROMO-10",
                    "Kiểm tra Báo cáo Hiệu quả Khuyến Mãi bên Admin (`/admin/promotions` -> Stats)",
                    "Đơn hàng có áp dụng mã `HPBD2026` đã được đặt thành công",
                    "1. Admin mở trang `/admin/promotions` -> Bấm tab 'Thống kê' (Stats).\n2. Kiểm tra số lượt đã dùng của mã `HPBD2026`.\n3. Kiểm tra tổng số tiền chiết khấu đã cấp và doanh thu đơn hàng mang lại.",
                    "Tab Thống kê Khuyến mãi",
                    "Số lượt dùng tăng lên đúng 1 lượt; Doanh thu và tiền giảm giá được tổng hợp chính xác vào báo cáo hiệu quả marketing.",
                    "Untested", "Major"
                )
            ]
        }
    ]

    # =========================================================================
    # CREATE THE 6 TESTER TABS (Each tester has 1 dedicated sheet!)
    # =========================================================================
    for tab_info in tabs_data:
        ws = wb.create_sheet(title=tab_info["sheet_name"])
        ws.views.sheetView[0].showGridLines = True
        
        # Row 1: Title Banner
        ws.merge_cells("A1:J1")
        title_cell = ws["A1"]
        title_cell.value = f"{tab_info['flow_title']} | {tab_info['tester_code']}: {tab_info['tester_name']} ({tab_info['persona']})"
        title_cell.font = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color=GOLD_BROWN, end_color=GOLD_BROWN, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 30

        # Row 2: Access & Credentials Info
        ws.merge_cells("A2:D2")
        ws["A2"].value = f"Tài khoản Test: {tab_info['account_info']}"
        ws["A2"].font = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
        ws["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells("E2:G2")
        ws["E2"].value = f"URL Truy cập: {tab_info['url_access']}"
        ws["E2"].font = Font(name=FONT_NAME, size=10, bold=False, color="0284C7")
        ws["E2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws["E2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws["H2"].value = "Tổng Test Cases:"
        ws["H2"].font = Font(name=FONT_NAME, size=10, bold=True, color="475569")
        ws["H2"].alignment = Alignment(horizontal="right", vertical="center")
        ws["H2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        ws["I2"].value = "=COUNTA(A4:A13)"
        ws["I2"].font = Font(name=FONT_NAME, size=11, bold=True, color=GOLD_BROWN)
        ws["I2"].alignment = Alignment(horizontal="center", vertical="center")
        ws["I2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

        ws["J2"].value = f"Mục tiêu: Đạt 100%"
        ws["J2"].font = Font(name=FONT_NAME, size=9, italic=True, color="64748B")
        ws["J2"].alignment = Alignment(horizontal="center", vertical="center")
        ws["J2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        ws.row_dimensions[2].height = 24

        # Row 3: Headers
        headers = [
            ("Mã TC", 14),
            ("Tên Kịch Bản", 28),
            ("Tiền Điều Kiện", 26),
            ("Các Bước Thực Hiện (Test Steps)", 45),
            ("Dữ Liệu Test", 25),
            ("Kết Quả Mong Đợi", 42),
            ("Kết Quả Thực Tế", 30),
            ("Trạng Thái", 15),
            ("Mức Độ", 14),
            ("Ghi Chú / Bug ID", 20)
        ]
        
        for col_idx, (hdr_text, col_width) in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_idx, value=hdr_text)
            cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = header_border
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = col_width
            
        ws.row_dimensions[3].height = 26

        # Row 4 to 13: Test Cases
        for row_offset, case_tuple in enumerate(tab_info["cases"]):
            curr_row = 4 + row_offset
            tc_code, tc_title, tc_pre, tc_steps, tc_data, tc_expected, tc_status, tc_severity = case_tuple
            
            row_data = [
                tc_code,
                tc_title,
                tc_pre,
                tc_steps,
                tc_data,
                tc_expected,
                "",             # Actual result (empty for tester to enter)
                tc_status,      # Untested
                tc_severity,    # Blocker / Critical / Major / Minor
                ""              # Bug ID
            ]
            
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=curr_row, column=col_idx, value=val)
                cell.font = Font(name=FONT_NAME, size=9.5)
                cell.border = thin_border
                
                # Alignments
                if col_idx in [1, 8, 9]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                elif col_idx in [4, 6, 7]:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    
                # Status styling
                if col_idx == 8:
                    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=STATUS_UNTESTED_FONT)
                    cell.fill = PatternFill(start_color=STATUS_UNTESTED_FILL, end_color=STATUS_UNTESTED_FILL, fill_type="solid")
                elif col_idx == 9:
                    sev_color = "991B1B" if val == "Blocker" else ("B45309" if val == "Critical" else ("1D4ED8" if val == "Major" else "475569"))
                    cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color=sev_color)
                elif col_idx == 1:
                    cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color=GOLD_BROWN)

            ws.row_dimensions[curr_row].height = 55

        # Freeze pane at row 4
        ws.freeze_panes = "A4"

    # =========================================================================
    # CREATE TAB 1: TỔNG QUAN & PHÂN CÔNG (Summary & Assignee Matrix)
    # =========================================================================
    # Remove default Sheet if exists, create or move to front
    default_sheet = wb["Sheet"] if "Sheet" in wb.sheetnames else None
    ws1 = wb.create_sheet(title="Tổng quan & Phân công", index=0)
    if default_sheet:
        wb.remove(default_sheet)
    ws1.views.sheetView[0].showGridLines = True

    # Main Title
    ws1.merge_cells("A1:K1")
    ws1["A1"].value = "KẾ HOẠCH & MA TRẬN PHÂN CÔNG KIỂM THỬ CHẤP NHẬN NGƯỜI DÙNG (UAT) - VELURA E-COMMERCE"
    ws1["A1"].font = Font(name=FONT_NAME, size=14, bold=True, color="FFFFFF")
    ws1["A1"].fill = PatternFill(start_color=GOLD_BROWN, end_color=GOLD_BROWN, fill_type="solid")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:K2")
    ws1["A2"].value = "Phiên bản Release Candidate: v1.0-RC | 6 Nhân Sự UAT - 6 Tab Phân Hệ End-to-End | Zero Formula Errors Standard"
    ws1["A2"].font = Font(name=FONT_NAME, size=10, italic=True, color="475569")
    ws1["A2"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 22

    # KPI Stat Cards (Rows 4 to 6)
    kpis = [
        ("B4", "C4", "B5", "C5", "B6", "C6", "TỔNG TEST CASES", 
         "='2.1 Mua Hàng E2E'!I2 + '2.2 Hủy Đơn & Đổi Trả E2E'!I2 + '2.3 AI Chatbot E2E'!I2 + '2.4 Quản Lý & Tra Cứu Đơn E2E'!I2 + '2.5 Dashboard & Phân Tích E2E'!I2 + '2.6 Khuyến Mãi & Đồng Bộ E2E'!I2", 
         "6 Quy trình Core E2E", "1E293B", "F8FAFC", "0F172A", False),
         
        ("E4", "F4", "E5", "F5", "E6", "F6", "ĐẠT (PASS)", 
         "=COUNTIF('2.1 Mua Hàng E2E'!H4:H13, \"Pass\") + COUNTIF('2.2 Hủy Đơn & Đổi Trả E2E'!H4:H13, \"Pass\") + COUNTIF('2.3 AI Chatbot E2E'!H4:H13, \"Pass\") + COUNTIF('2.4 Quản Lý & Tra Cứu Đơn E2E'!H4:H13, \"Pass\") + COUNTIF('2.5 Dashboard & Phân Tích E2E'!H4:H13, \"Pass\") + COUNTIF('2.6 Khuyến Mãi & Đồng Bộ E2E'!H4:H13, \"Pass\")", 
         "Kịch bản vượt qua", "166534", "DCFCE7", "166534", False),
         
        ("H4", "I4", "H5", "I5", "H6", "I6", "LỖI (FAIL)", 
         "=COUNTIF('2.1 Mua Hàng E2E'!H4:H13, \"Fail\") + COUNTIF('2.2 Hủy Đơn & Đổi Trả E2E'!H4:H13, \"Fail\") + COUNTIF('2.3 AI Chatbot E2E'!H4:H13, \"Fail\") + COUNTIF('2.4 Quản Lý & Tra Cứu Đơn E2E'!H4:H13, \"Fail\") + COUNTIF('2.5 Dashboard & Phân Tích E2E'!H4:H13, \"Fail\") + COUNTIF('2.6 Khuyến Mãi & Đồng Bộ E2E'!H4:H13, \"Fail\")", 
         "Cần DEV xử lý", "991B1B", "FEE2E2", "991B1B", False),
         
        ("J4", "K4", "J5", "K5", "J6", "K6", "TỶ LỆ ĐẠT (PASS RATE)", 
         "=IF(B5>0, E5/B5, 0)", 
         "Mục tiêu: >= 95%", "7D562D", "FEF3C7", "B45309", True)
    ]

    for c_start, c_end, val_start, val_end, sub_start, sub_end, title, formula, subtitle, text_col, bg_col, num_col, is_pct in kpis:
        ws1.merge_cells(f"{c_start}:{c_end}")
        ws1[c_start].value = title
        ws1[c_start].font = Font(name=FONT_NAME, size=9.5, bold=True, color=text_col)
        ws1[c_start].fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
        ws1[c_start].alignment = Alignment(horizontal="center", vertical="center")

        ws1.merge_cells(f"{val_start}:{val_end}")
        ws1[val_start].value = formula
        ws1[val_start].font = Font(name=FONT_NAME, size=18, bold=True, color=num_col)
        ws1[val_start].fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
        ws1[val_start].alignment = Alignment(horizontal="center", vertical="center")
        if is_pct:
            ws1[val_start].number_format = '0.0%'

        ws1.merge_cells(f"{sub_start}:{sub_end}")
        ws1[sub_start].value = subtitle
        ws1[sub_start].font = Font(name=FONT_NAME, size=8.5, italic=True, color="64748B")
        ws1[sub_start].fill = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")
        ws1[sub_start].alignment = Alignment(horizontal="center", vertical="center")

    ws1.row_dimensions[4].height = 20
    ws1.row_dimensions[5].height = 32
    ws1.row_dimensions[6].height = 18

    # Section 1: Tester Assignment Table
    ws1.cell(row=8, column=1, value="1. MA TRẬN PHÂN CÔNG 6 NHÂN SỰ UAT & TÀI KHOẢN TRUY CẬP (MỖI NGƯỜI 1 TAB)").font = Font(name=FONT_NAME, size=11, bold=True, color="1E293B")
    
    assign_headers = [
        ("Mã Tester", 12),
        ("Họ và Tên", 18),
        ("Vai Trò (Persona)", 25),
        ("Tab Quy Trình Phụ Trách", 25),
        ("Tài Khoản Đăng Nhập", 26),
        ("Mật Khẩu Mặc Định", 16),
        ("Số Lượng TC", 14),
        ("Số TC Đạt", 12),
        ("Số TC Lỗi", 12),
        ("Tỷ Lệ Đạt", 14),
        ("Trạng Thái Tiến Độ", 18)
    ]
    for col_idx, (h_text, h_width) in enumerate(assign_headers, 1):
        cell = ws1.cell(row=9, column=col_idx, value=h_text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws1.column_dimensions[col_letter].width = h_width
    ws1.row_dimensions[9].height = 26

    testers = [
        ("Tester 01", "Nguyễn Thị Ánh", "Khách Mua Sắm Online", "2.1 Mua Hàng E2E", "shopper.uat@velura.test", "Velura@2026", "='2.1 Mua Hàng E2E'!I2", "=COUNTIF('2.1 Mua Hàng E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.1 Mua Hàng E2E'!H4:H13, \"Fail\")", "=IF(G10>0, H10/G10, 0)"),
        ("Tester 02", "Trần Quốc Bảo", "Khách Hậu Mãi & CSKH RMA", "2.2 Hủy Đơn & Đổi Trả E2E", "customer.rma@velura.test", "Velura@2026", "='2.2 Hủy Đơn & Đổi Trả E2E'!I2", "=COUNTIF('2.2 Hủy Đơn & Đổi Trả E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.2 Hủy Đơn & Đổi Trả E2E'!H4:H13, \"Fail\")", "=IF(G11>0, H11/G11, 0)"),
        ("Tester 03", "Lê Hoàng Châu", "Trải Nghiệm AI Stylist Bot", "2.3 AI Chatbot E2E", "member.chatbot@velura.test", "Velura@2026", "='2.3 AI Chatbot E2E'!I2", "=COUNTIF('2.3 AI Chatbot E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.3 AI Chatbot E2E'!H4:H13, \"Fail\")", "=IF(G12>0, H12/G12, 0)"),
        ("Tester 04", "Phạm Văn Dũng", "Tra Cứu Đơn & Kho Fulfillment", "2.4 Quản Lý & Tra Cứu Đơn E2E", "warehouse.ops@velura.vn", "Admin@2026", "='2.4 Quản Lý & Tra Cứu Đơn E2E'!I2", "=COUNTIF('2.4 Quản Lý & Tra Cứu Đơn E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.4 Quản Lý & Tra Cứu Đơn E2E'!H4:H13, \"Fail\")", "=IF(G13>0, H13/G13, 0)"),
        ("Tester 05", "Đặng Thị Mai", "Phân Tích Dữ Liệu & Giám Đốc", "2.5 Dashboard & Phân Tích E2E", "admin.analyst@velura.vn", "Admin@2026", "='2.5 Dashboard & Phân Tích E2E'!I2", "=COUNTIF('2.5 Dashboard & Phân Tích E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.5 Dashboard & Phân Tích E2E'!H4:H13, \"Fail\")", "=IF(G14>0, H14/G14, 0)"),
        ("Tester 06", "Vũ Minh Tuấn", "Trưởng Phòng MKT Khuyến Mãi", "2.6 Khuyến Mãi & Đồng Bộ E2E", "admin.marketing@velura.vn", "Admin@2026", "='2.6 Khuyến Mãi & Đồng Bộ E2E'!I2", "=COUNTIF('2.6 Khuyến Mãi & Đồng Bộ E2E'!H4:H13, \"Pass\")", "=COUNTIF('2.6 Khuyến Mãi & Đồng Bộ E2E'!H4:H13, \"Fail\")", "=IF(G15>0, H15/G15, 0)")
    ]

    for idx, t in enumerate(testers, 10):
        t_code, t_name, t_role, t_tab, t_user, t_pass, f_total, f_pass, f_fail, f_rate = t
        ws1.cell(row=idx, column=1, value=t_code).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=idx, column=1).font = Font(name=FONT_NAME, size=10, bold=True, color=GOLD_BROWN)
        ws1.cell(row=idx, column=2, value=t_name).font = Font(name=FONT_NAME, size=10, bold=True)
        ws1.cell(row=idx, column=3, value=t_role).font = Font(name=FONT_NAME, size=9.5)
        ws1.cell(row=idx, column=4, value=t_tab).font = Font(name=FONT_NAME, size=9.5, bold=True, color=ACCENT_BLUE)
        ws1.cell(row=idx, column=5, value=t_user).font = Font(name=FONT_NAME, size=9.5)
        ws1.cell(row=idx, column=6, value=t_pass).font = Font(name=FONT_NAME, size=9.5, color="475569")
        
        ws1.cell(row=idx, column=7, value=f_total).font = Font(name=FONT_NAME, size=10, bold=True)
        ws1.cell(row=idx, column=7).alignment = Alignment(horizontal="center", vertical="center")

        ws1.cell(row=idx, column=8, value=f_pass).font = Font(name=FONT_NAME, size=10, bold=True, color="166534")
        ws1.cell(row=idx, column=8).alignment = Alignment(horizontal="center", vertical="center")

        ws1.cell(row=idx, column=9, value=f_fail).font = Font(name=FONT_NAME, size=10, bold=True, color="991B1B")
        ws1.cell(row=idx, column=9).alignment = Alignment(horizontal="center", vertical="center")

        ws1.cell(row=idx, column=10, value=f_rate).font = Font(name=FONT_NAME, size=10, bold=True)
        ws1.cell(row=idx, column=10).alignment = Alignment(horizontal="center", vertical="center")
        ws1.cell(row=idx, column=10).number_format = '0.0%'

        ws1.cell(row=idx, column=11, value=f'=IF(J{idx}>=0.95, "Sẵn sàng Release", IF(I{idx}>0, "Đang có Bug chặn", "Đang kiểm thử"))')
        ws1.cell(row=idx, column=11).font = Font(name=FONT_NAME, size=9.5, italic=True)
        ws1.cell(row=idx, column=11).alignment = Alignment(horizontal="center", vertical="center")

        for c in range(1, 12):
            ws1.cell(row=idx, column=c).border = thin_border
        ws1.row_dimensions[idx].height = 24

    # Summary Row 16
    ws1.merge_cells("A16:F16")
    ws1["A16"].value = "TỔNG CỘNG HỆ THỐNG"
    ws1["A16"].font = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
    ws1["A16"].alignment = Alignment(horizontal="right", vertical="center")
    ws1["A16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws1["G16"].value = "=SUM(G10:G15)"
    ws1["G16"].font = Font(name=FONT_NAME, size=11, bold=True, color=GOLD_BROWN)
    ws1["G16"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["G16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws1["H16"].value = "=SUM(H10:H15)"
    ws1["H16"].font = Font(name=FONT_NAME, size=11, bold=True, color="166534")
    ws1["H16"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["H16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws1["I16"].value = "=SUM(I10:I15)"
    ws1["I16"].font = Font(name=FONT_NAME, size=11, bold=True, color="991B1B")
    ws1["I16"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["I16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    ws1["J16"].value = "=IF(G16>0, H16/G16, 0)"
    ws1["J16"].font = Font(name=FONT_NAME, size=11, bold=True, color=GOLD_BROWN)
    ws1["J16"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["J16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws1["J16"].number_format = '0.0%'

    ws1["K16"].value = '=IF(J16>=0.95, "ĐẠT NGHIỆM THU", "CHƯA ĐẠT")'
    ws1["K16"].font = Font(name=FONT_NAME, size=10, bold=True, color="1E293B")
    ws1["K16"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["K16"].fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    for c in range(1, 12):
        ws1.cell(row=16, column=c).border = thin_border
    ws1.row_dimensions[16].height = 26

    # Section 2: Severity Standard & SLA
    ws1.cell(row=18, column=1, value="2. QUY CHUẨN ĐÁNH GIÁ MỨC ĐỘ LỖI (DEFECT SEVERITY) & CAM KẾT XỬ LÝ (SLA)").font = Font(name=FONT_NAME, size=11, bold=True, color="1E293B")
    
    sev_headers = [("Mức độ", 12), ("Định nghĩa kỹ thuật", 35), ("Tác động nghiệp vụ", 28), ("Ví dụ thực tế trong Velura", 35), ("Cam kết SLA Fix", 18)]
    for c_idx, (sh_text, sh_w) in enumerate(sev_headers, 1):
        cell = ws1.cell(row=19, column=c_idx, value=sh_text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
    ws1.row_dimensions[19].height = 24

    sevs = [
        ("Blocker", "Lỗi sập luồng (Crash/500), dừng toàn bộ quy trình, không có cách giải quyết thay thế (No Workaround).", "Khách không thể thanh toán, mất doanh thu lập tức.", "Lỗi thanh toán Stripe 500; Hết hàng tồn nhưng nút Mua ngay vẫn không khóa; Kho không nhận được hàng.", "Khắc phục <= 2h"),
        ("Critical", "Chức năng chính bị hỏng nặng, sai lệch tiền tệ, tính sai giá trị hoàn trả hoặc sai tồn kho.", "Ảnh hưởng dữ liệu tài chính, thất thoát doanh thu hoặc tiền hoàn.", "Số tiền hoàn đổi trả cho phép sửa tay; Voucher áp dụng sai chiết khấu; Đơn hàng không cập nhật vào Dashboard.", "Khắc phục <= 6h"),
        ("Major", "Chức năng quan trọng không đúng thiết kế nhưng có phương án giải quyết tạm thời.", "Gây trải nghiệm xấu cho người dùng, làm gián đoạn nhưng vẫn hoàn tất được đơn.", "Dropdown menu thao tác đổi trả bị vỡ giao diện; Chatbot phản hồi chậm > 10s; Voucher sinh nhật không tự nạp.", "Khắc phục <= 24h"),
        ("Minor", "Lỗi giao diện (UI), sai chính tả, lệch khoảng cách padding, thiếu tooltip thông tin.", "Không ảnh hưởng đến logic nghiệp vụ cốt lõi.", "Icon căn lệch 2px; Màu sắc badge chưa chuẩn token thiết kế; Thiếu placeholder hướng dẫn.", "Khắc phục Sprint tới")
    ]

    for s_idx, (s_name, s_def, s_imp, s_ex, s_sla) in enumerate(sevs, 20):
        c1 = ws1.cell(row=s_idx, column=1, value=s_name)
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name=FONT_NAME, size=10, bold=True, 
                       color="991B1B" if s_name=="Blocker" else ("B45309" if s_name=="Critical" else ("1D4ED8" if s_name=="Major" else "475569")))
        
        ws1.cell(row=s_idx, column=2, value=s_def).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=s_idx, column=3, value=s_imp).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws1.cell(row=s_idx, column=4, value=s_ex).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        c5 = ws1.cell(row=s_idx, column=5, value=s_sla)
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.font = Font(name=FONT_NAME, size=9.5, bold=True, color="1E293B")

        for c in range(1, 6):
            ws1.cell(row=s_idx, column=c).border = thin_border
            ws1.cell(row=s_idx, column=c).font = Font(name=FONT_NAME, size=9.5)
            if c == 1:
                ws1.cell(row=s_idx, column=c).font = Font(name=FONT_NAME, size=10, bold=True, color="991B1B" if s_name=="Blocker" else ("B45309" if s_name=="Critical" else ("1D4ED8" if s_name=="Major" else "475569")))
        ws1.row_dimensions[s_idx].height = 36

    # =========================================================================
    # CREATE TAB 8: NHẬT KÝ BÁO CÁO BUG UAT (Defect Tracking Sheet)
    # =========================================================================
    ws_bug = wb.create_sheet(title="Nhật ký Báo cáo Bug UAT")
    ws_bug.views.sheetView[0].showGridLines = True

    ws_bug.merge_cells("A1:K1")
    ws_bug["A1"].value = "NHẬT KÝ THEO DÕI VÀ BÁO CÁO LỖI PHÁT SINH TRONG QUÁ TRÌNH UAT (DEFECT LOG)"
    ws_bug["A1"].font = Font(name=FONT_NAME, size=12, bold=True, color="FFFFFF")
    ws_bug["A1"].fill = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    ws_bug["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_bug.row_dimensions[1].height = 32

    bug_headers = [
        ("Bug ID", 12),
        ("Quy Trình / Tab", 22),
        ("Mã TC Liên Quan", 14),
        ("Tóm Tắt Lỗi (Defect Summary)", 38),
        ("Mức Độ", 12),
        ("Các Bước Tái Hiện (Steps to Reproduce)", 40),
        ("Kết Quả Thực Tế Bị Lỗi", 32),
        ("Người Báo Lỗi", 16),
        ("Trạng Thái Fix", 16),
        ("Dev Phụ Trách", 16),
        ("Ghi Chú / Link PR Fix", 22)
    ]

    for col_idx, (bh_text, bh_width) in enumerate(bug_headers, 1):
        cell = ws_bug.cell(row=2, column=col_idx, value=bh_text)
        cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = header_border
        col_letter = get_column_letter(col_idx)
        ws_bug.column_dimensions[col_letter].width = bh_width
    ws_bug.row_dimensions[2].height = 26

    sample_bugs = [
        ("BUG-001", "2.1 Mua Hàng E2E", "TC-BUY-04", "Biến thể hết hàng tồn kho nhưng nút Mua ngay không bị vô hiệu hóa", "Blocker", "1. Chọn sản phẩm SP-001.\n2. Chọn Màu Đen, Size XL (Stock = 0).\n3. Nút Mua ngay vẫn bấm được và chuyển sang Checkout.", "Cho phép đặt hàng sản phẩm không còn trong kho dẫn đến âm tồn kho.", "Nguyễn Thị Ánh", "Fixed", "Dev Core", "PR #142 (Closed)"),
        ("BUG-002", "2.2 Hủy Đơn & Đổi Trả E2E", "TC-RMA-05", "Số tiền hoàn đổi trả cho phép sửa tay dẫn đến nguy cơ sai lệch kế toán", "Critical", "1. Mở form yêu cầu đổi trả `/account/returns`.\n2. Chọn hình thức hoàn tiền.\n3. Ô số tiền cho phép gõ số tiền lớn hơn hóa đơn gốc.", "Khách hàng có thể gian lận sửa số tiền hoàn cao hơn thực tế.", "Trần Quốc Bảo", "Fixed", "Dev Core", "PR #145 (Closed)"),
        ("BUG-003", "2.3 AI Chatbot E2E", "TC-BOT-04", "Nút Thêm vào giỏ trên Card sản phẩm Chatbot không cập nhật badge giỏ", "Major", "1. Chatbot gợi ý sản phẩm.\n2. Bấm nút Thêm vào giỏ trên card.\n3. Header giỏ hàng không nhảy số ngay mà phải F5 trang.", "Trải nghiệm không mượt mà, khách tưởng bấm trượt nên bấm nhiều lần.", "Lê Hoàng Châu", "Open", "Dev AI", "Jira KAN-48"),
        ("BUG-004", "2.6 Khuyến Mãi & Đồng Bộ E2E", "TC-PROMO-08", "Voucher không kiểm tra điều kiện giá trị đơn hàng tối thiểu tại Checkout", "Critical", "1. Giỏ hàng 300.000đ.\n2. Nhập mã VELURA20 (yêu cầu min 500k).\n3. Hệ thống vẫn trừ tiền 20% bình thường.", "Sai lệch chính sách marketing, vi phạm biên lợi nhuận công ty.", "Vũ Minh Tuấn", "Fixed", "Dev Core", "PR #148 (Closed)")
    ]

    for b_idx, bug_row in enumerate(sample_bugs, 3):
        for col_idx, b_val in enumerate(bug_row, 1):
            cell = ws_bug.cell(row=b_idx, column=col_idx, value=b_val)
            cell.font = Font(name=FONT_NAME, size=9.5)
            cell.border = thin_border
            if col_idx in [1, 3, 5, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if col_idx == 1:
                cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color="991B1B")
            elif col_idx == 5:
                cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color="991B1B" if b_val=="Blocker" else "B45309")
            elif col_idx == 9:
                cell.font = Font(name=FONT_NAME, size=9.5, bold=True, color="166534" if b_val=="Fixed" else "B45309")
        ws_bug.row_dimensions[b_idx].height = 42

    # Blank rows for logging new defects
    for empty_idx in range(len(sample_bugs) + 3, len(sample_bugs) + 20):
        b_code = f"BUG-{empty_idx-2:03d}"
        c = ws_bug.cell(row=empty_idx, column=1, value=b_code)
        c.font = Font(name=FONT_NAME, size=9.5, color="94A3B8")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
        for col_idx in range(2, 12):
            ws_bug.cell(row=empty_idx, column=col_idx).border = thin_border
        ws_bug.row_dimensions[empty_idx].height = 24

    ws_bug.freeze_panes = "A3"

    # Save workbook
    output_dir = os.path.join("docs", "ba")
    os.makedirs(output_dir, exist_ok=True)
    primary_file = os.path.join(output_dir, "Velura_UAT_E2E_Workflows_6_Testers.xlsx")
    wb.save(primary_file)
    print(f"UAT Test Plan workbook with 6 dedicated tabs created successfully at: {primary_file}")
    
    # Try saving to the previous name if not locked by Excel
    alt_file = os.path.join(output_dir, "Velura_UAT_Test_Plan_Core_Workflows_6_Testers.xlsx")
    try:
        wb.save(alt_file)
        print(f"Also updated: {alt_file}")
    except PermissionError:
        print(f"Note: {alt_file} is currently open in Excel. Kept updated file at: {primary_file}")

if __name__ == "__main__":
    build_uat_workbook()
